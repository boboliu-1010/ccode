from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUTLINE = ROOT / 'docs/zh/10-Claude Code 源码深度解读-分享提纲.md'
OUT = ROOT / 'docs/zh/10-Claude Code 源码调研.pptx'

FONT = 'Arial'
MONO = 'Courier New'

BG = RGBColor(4, 10, 24)
BG_2 = RGBColor(7, 16, 35)
PANEL = RGBColor(10, 22, 46)
PANEL_2 = RGBColor(14, 31, 63)
PANEL_3 = RGBColor(12, 26, 54)
CODE_BG = RGBColor(7, 15, 31)
TITLE = RGBColor(245, 248, 255)
TEXT = RGBColor(222, 230, 244)
MUTED = RGBColor(153, 171, 196)
ACCENT = RGBColor(72, 198, 255)
ACCENT_2 = RGBColor(156, 163, 255)
ACCENT_3 = RGBColor(61, 224, 199)
WARN = RGBColor(250, 204, 21)
LINE = RGBColor(38, 63, 110)
CODE_TEXT = RGBColor(190, 246, 255)


class SlideData:
    def __init__(self, num: int, title: str, fields: Dict[str, str]):
        self.num = num
        self.title = title
        self.fields = fields


def strip_links(text: str) -> str:
    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
    return text.replace('`', '').strip()


def parse_bullets(block: str) -> List[str]:
    items: List[str] = []
    for line in block.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith('- '):
            items.append(strip_links(s[2:]))
        elif re.match(r'^\d+\.\s+', s):
            items.append(strip_links(re.sub(r'^\d+\.\s+', '', s)))
    return items


def parse_code(block: str) -> str | None:
    m = re.search(r'```(?:\w+)?\n(.*?)```', block, re.S)
    if not m:
        return None
    return m.group(1).strip()


def parse_fields(block: str) -> Dict[str, str]:
    fields: Dict[str, str] = {}
    current = None
    buf: List[str] = []
    for line in block.splitlines():
        m = re.match(r'^\*\*(.+?)\*\*$', line.strip())
        if m:
            if current is not None:
                fields[current] = '\n'.join(buf).strip()
            current = m.group(1)
            buf = []
        else:
            if current is not None:
                buf.append(line)
    if current is not None:
        fields[current] = '\n'.join(buf).strip()
    return fields


def parse_outline() -> List[SlideData]:
    text = OUTLINE.read_text()
    chunks = re.split(r'(?=^## 第 \d+ 页：)', text, flags=re.M)
    slides: List[SlideData] = []
    for chunk in chunks:
        if not chunk.startswith('## 第 '):
            continue
        lines = chunk.splitlines()
        m = re.match(r'^## 第 (\d+) 页：(.+)$', lines[0])
        if not m:
            continue
        num = int(m.group(1))
        title = m.group(2).strip()
        fields = parse_fields('\n'.join(lines[1:]))
        slides.append(SlideData(num, title, fields))
    return slides


def set_run_font(run, size, color, bold=False, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def style_tf(tf, top=4, bottom=4, left=8, right=8):
    tf.word_wrap = True
    tf.margin_top = Pt(top)
    tf.margin_bottom = Pt(bottom)
    tf.margin_left = Pt(left)
    tf.margin_right = Pt(right)
    tf.vertical_anchor = MSO_ANCHOR.TOP


def add_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(12.1), Inches(0.7))
    tf = tb.text_frame
    style_tf(tf, 0, 0, 0, 0)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run_font(r, 24, TITLE, True)
    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.58), Inches(0.92), Inches(12.0), Inches(0.38))
        tf = sb.text_frame
        style_tf(tf, 0, 0, 0, 0)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = subtitle
        set_run_font(r, 11.5, MUTED)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.58), Inches(1.28), Inches(2.4), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def add_shape(slide, kind, x, y, w, h, fill=PANEL, line=LINE):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    return shp


def add_panel(slide, x, y, w, h, title, title_color=ACCENT, fill=PANEL):
    body = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill)
    hdr = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.32))
    hdr.fill.solid()
    hdr.fill.fore_color.rgb = PANEL_2
    hdr.line.fill.background()
    tf = hdr.text_frame
    style_tf(tf, 2, 2, 5, 5)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run_font(r, 11, title_color, True)
    return body


def add_text(slide, x, y, w, h, lines: List[str], size=16, color=TEXT, bullet=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    style_tf(tf)
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.bullet = bullet
        p.level = 0
        p.space_after = Pt(5)
        r = p.add_run()
        r.text = line
        set_run_font(r, size, color)
    return tb


def add_rich_lines(slide, x, y, w, h, sections: List[Tuple[str, List[str]]]):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    style_tf(tf)
    tf.clear()
    first = True
    for heading, items in sections:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(5)
        p.bullet = False
        r = p.add_run()
        r.text = heading
        set_run_font(r, 14.5, ACCENT_2, True)
        for item in items:
            pp = tf.add_paragraph()
            pp.bullet = True
            pp.space_after = Pt(4)
            rr = pp.add_run()
            rr.text = item
            set_run_font(rr, 13.4, TEXT)
    return tb


def add_code_box(slide, x, y, w, h, code, label='关键源码片段'):
    body = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill=CODE_BG, line=ACCENT)
    hdr = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.28))
    hdr.fill.solid()
    hdr.fill.fore_color.rgb = PANEL_2
    hdr.line.fill.background()
    tf = hdr.text_frame
    style_tf(tf, 2, 2, 6, 6)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = label
    set_run_font(r, 10, ACCENT, True, MONO)
    tf = body.text_frame
    style_tf(tf, 14, 6, 10, 10)
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = code
    set_run_font(r, 10.5, CODE_TEXT, False, MONO)
    return body


def add_footer(slide, left, right=''):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.1), Inches(0.2))
    tf = tb.text_frame
    style_tf(tf, 0, 0, 0, 0)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = left
    set_run_font(r, 10, MUTED)
    if right:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.RIGHT
        r2 = p2.add_run()
        r2.text = right
        set_run_font(r2, 9.5, MUTED)


def markdown_label_list(block: str) -> List[str]:
    items = parse_bullets(block)
    labels = []
    for item in items:
        labels.append(strip_links(item))
    return labels


def connect(slide, x1, y1, x2, y2, color=ACCENT):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(1.6)
    line.line.end_arrowhead = True
    return line


def add_cover_slide(prs: Presentation, sd: SlideData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG_2)
    add_title(slide, 'Claude Code 源码调研', '总体架构、工作流程、功能设计、产出形态与开发流程建议')
    hero = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.62, 1.72, 12.0, 0.88, fill=RGBColor(8, 24, 52), line=ACCENT)
    tf = hero.text_frame
    style_tf(tf, 6, 6, 12, 12)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = '通过足够的源码理解，帮助用户和团队更好地使用 Claude Code。'
    set_run_font(r, 19, TITLE, True)
    add_panel(slide, 0.72, 2.95, 5.72, 2.5, '这场分享要解决什么', title_color=ACCENT)
    add_text(slide, 0.95, 3.32, 5.22, 2.0, [
        '为什么同样是 Claude Code，有的人越用越顺，有的人越用越乱。',
        '为什么它更适合工程任务，而不是自由聊天式请求。',
        '怎样组织任务、边界和验证，才能让它稳定地产出结果。',
    ], size=16.5)
    add_panel(slide, 6.74, 2.95, 5.86, 2.5, '这场分享怎么讲', title_color=ACCENT_2)
    add_text(slide, 6.97, 3.32, 5.34, 2.0, [
        '先用源码建立最小必要架构模型。',
        '再讲工作流程、功能和典型产出。',
        '最后把结论落到开发流程建议与代码走读入口。',
    ], size=16.5)
    add_panel(slide, 0.72, 5.82, 11.88, 0.9, '核心判断', title_color=WARN, fill=PANEL_3)
    add_text(slide, 0.93, 6.17, 11.45, 0.35, ['Claude Code 不是自由聊天助手，而是一套围绕工程任务构建的 terminal agent runtime（终端代理运行时）。'], size=17, bullet=False)
    add_footer(slide, 'Claude Code 源码调研', '封面')


def add_overview_diagram(prs: Presentation, sd: SlideData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, sd.title, '从用户输入到主执行链，再到控制面和续航能力。')
    q = parse_bullets(sd.fields.get('这一页要回答的问题', ''))
    if q:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.62, 1.55, 12.02, 0.58, fill=PANEL_3, line=ACCENT)
        add_text(slide, 0.82, 1.72, 11.6, 0.22, q, size=14.5, bullet=False)
    boxes = [
        ('用户输入', 0.78, 2.55, 1.55, 0.7, ACCENT),
        ('Prompt Stack\n提示词栈', 2.6, 2.48, 1.95, 0.84, ACCENT),
        ('QueryEngine\n会话宿主', 4.95, 2.48, 1.95, 0.84, ACCENT_2),
        ('query.ts\nTurn Loop', 7.28, 2.48, 1.95, 0.84, ACCENT_2),
        ('Tool Pipeline\n工具执行链', 9.62, 2.44, 2.35, 0.92, ACCENT_3),
    ]
    for txt, x, y, w, h, color in boxes:
        shp = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill=PANEL, line=color)
        tf = shp.text_frame
        style_tf(tf, 8, 8, 8, 8)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        for idx, line in enumerate(txt.split('\n')):
            r = p.add_run()
            r.text = ('' if idx == 0 else '\n') + line
            set_run_font(r, 15 if idx == 0 else 12.5, TITLE if idx == 0 else color, idx == 0)
    connect(slide, 2.33, 2.9, 2.6, 2.9)
    connect(slide, 4.55, 2.9, 4.95, 2.9)
    connect(slide, 6.9, 2.9, 7.28, 2.9)
    connect(slide, 9.24, 2.9, 9.62, 2.9)
    cp = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.92, 4.38, 4.2, 1.0, fill=PANEL, line=WARN)
    tf = cp.text_frame
    style_tf(tf, 6, 6, 10, 10)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = 'Settings / Auth / Policy / Prompt\n控制面'
    set_run_font(r, 16, TITLE, True)
    obs = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 5.42, 4.38, 3.28, 1.0, fill=PANEL, line=ACCENT_2)
    tf = obs.text_frame
    style_tf(tf, 6, 6, 10, 10)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = 'Transcript / Recovery\nCompact / Observability'
    set_run_font(r, 15.5, TITLE, True)
    env = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 9.02, 4.38, 3.22, 1.0, fill=PANEL, line=ACCENT_3)
    tf = env.text_frame
    style_tf(tf, 6, 6, 10, 10)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = 'Tools / Tasks / MCP\n执行环境与扩展能力'
    set_run_font(r, 15.5, TITLE, True)
    connect(slide, 5.95, 3.32, 5.95, 4.38, WARN)
    connect(slide, 8.25, 3.32, 7.1, 4.38, ACCENT_2)
    connect(slide, 10.82, 3.36, 10.82, 4.38, ACCENT_3)
    takeaway = parse_bullets(sd.fields.get('希望听众带走什么', ''))
    if takeaway:
        add_panel(slide, 0.82, 6.08, 11.85, 0.72, '这一页想说明什么', title_color=ACCENT)
        add_text(slide, 1.04, 6.38, 11.35, 0.2, takeaway, size=14.5, bullet=False)
    add_footer(slide, 'Claude Code 源码调研', '总体架构')


def add_workflow_diagram(prs: Presentation, sd: SlideData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, sd.title, '一次请求如何被推进、约束、回灌、压缩和恢复。')
    q = parse_bullets(sd.fields.get('这一页要回答的问题', ''))
    if q:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.62, 1.55, 12.02, 0.58, fill=PANEL_3, line=ACCENT)
        add_text(slide, 0.82, 1.72, 11.6, 0.22, q, size=14.5, bullet=False)
    steps = [
        ('用户输入', 0.82, 2.55, 1.45),
        ('装配环境', 2.45, 2.55, 1.5),
        ('会话宿主', 4.18, 2.55, 1.5),
        ('turn loop', 5.92, 2.55, 1.5),
        ('tool pipeline', 7.65, 2.55, 1.8),
        ('结果回灌', 9.7, 2.55, 1.5),
        ('继续 / compact / recovery', 4.8, 4.45, 3.4),
    ]
    colors = [ACCENT, ACCENT, ACCENT_2, ACCENT_2, ACCENT_3, ACCENT_3, WARN]
    for (txt, x, y, w), color in zip(steps, colors):
        shp = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, 0.78, fill=PANEL, line=color)
        tf = shp.text_frame
        style_tf(tf, 8, 8, 8, 8)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = txt
        set_run_font(r, 15, TITLE, True)
    connect(slide, 2.27, 2.94, 2.45, 2.94)
    connect(slide, 3.95, 2.94, 4.18, 2.94)
    connect(slide, 5.68, 2.94, 5.92, 2.94)
    connect(slide, 7.42, 2.94, 7.65, 2.94)
    connect(slide, 9.45, 2.94, 9.7, 2.94)
    connect(slide, 10.45, 3.33, 8.15, 4.45, WARN)
    connect(slide, 6.5, 4.45, 6.5, 3.33, WARN)
    core = parse_bullets(sd.fields.get('核心内容', ''))
    add_panel(slide, 0.86, 5.62, 11.8, 1.05, '流程要点', title_color=ACCENT_2)
    add_text(slide, 1.06, 5.98, 11.35, 0.55, core[:4], size=13.8)
    add_footer(slide, 'Claude Code 源码调研', '工作流程')


def add_capability_slide(prs: Presentation, sd: SlideData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, sd.title, '先从用户感知的能力面理解，再进入主要功能。')
    thesis = parse_bullets(sd.fields.get('这一页要回答的问题', ''))
    if thesis:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.62, 1.55, 12.02, 0.58, fill=PANEL_3, line=ACCENT)
        add_text(slide, 0.82, 1.72, 11.6, 0.22, thesis, size=14.5, bullet=False)
    cards = [
        ('交互', ['REPL', 'commands', 'structured output'], 0.82, 2.5, ACCENT),
        ('执行', ['Read / Edit / Bash', 'tool pipeline', 'task runtime'], 3.05, 2.5, ACCENT_2),
        ('扩展', ['skills', 'MCP', 'plugins', 'remote capability'], 5.28, 2.5, ACCENT_3),
        ('控制', ['prompt stack', 'settings', 'policy', 'permissions'], 7.51, 2.5, WARN),
        ('续航', ['transcript', 'recovery', 'compact', 'replacement'], 9.74, 2.5, ACCENT),
        ('可观测性', ['profiler', 'cache break', 'context analysis'], 3.95, 4.75, ACCENT_2),
    ]
    for title, items, x, y, color in cards:
        add_panel(slide, x, y, 2.05, 1.6, title, title_color=color)
        add_text(slide, x+0.08, y+0.37, 1.9, 1.1, items, size=12.6)
    takeaway = parse_bullets(sd.fields.get('希望听众带走什么', ''))
    add_panel(slide, 0.86, 6.15, 11.8, 0.65, '这一页想说明什么', title_color=ACCENT_3)
    add_text(slide, 1.06, 6.42, 11.35, 0.2, takeaway or ['功能多本身不重要，关键是这些能力能围绕主执行链协同。'], size=14.2, bullet=False)
    add_footer(slide, 'Claude Code 源码调研', '功能介绍')


def generic_slide(prs: Presentation, sd: SlideData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    subtitle = ''
    if 1 <= sd.num <= 8:
        subtitle = '总体架构'
    elif 9 <= sd.num <= 16:
        subtitle = '工作流程'
    elif 17 <= sd.num <= 35:
        subtitle = '功能介绍'
    elif 36 <= sd.num <= 40:
        subtitle = '产出'
    elif 41 <= sd.num <= 53:
        subtitle = '开发流程建议'
    else:
        subtitle = '代码走读概略'
    add_title(slide, sd.title, subtitle)
    question = parse_bullets(sd.fields.get('这一页要回答的问题', ''))
    core = parse_bullets(sd.fields.get('核心内容', ''))
    sources = markdown_label_list(sd.fields.get('代码支撑', ''))
    insight = parse_bullets(sd.fields.get('代码理解支撑', ''))
    takeaway = parse_bullets(sd.fields.get('希望听众带走什么', ''))
    code = parse_code(sd.fields.get('关键代码片段', ''))

    if question:
        add_panel(slide, 0.62, 1.5, 12.02, 0.78, '这一页要回答的问题', title_color=ACCENT)
        add_text(slide, 0.82, 1.85, 11.6, 0.25, question, size=14.5, bullet=False)
        top_y = 2.48
    else:
        top_y = 1.55

    add_panel(slide, 0.62, top_y, 7.0, 3.95, '核心内容', title_color=ACCENT_2)
    if core:
        add_text(slide, 0.82, top_y+0.38, 6.65, 3.48, core[:9], size=15.2)
    else:
        sections = []
        for key in ('功能', '实现原理', '使用技巧'):
            if key in sd.fields:
                sections.append((key, parse_bullets(sd.fields.get(key, '')) or [strip_links(sd.fields.get(key, ''))]))
        add_rich_lines(slide, 0.8, top_y+0.38, 6.68, 3.42, sections)

    if code:
        add_code_box(slide, 7.82, top_y, 4.8, 1.95, code, '关键源码片段')
        add_panel(slide, 7.82, top_y+2.12, 4.8, 1.83, '这段代码说明什么', title_color=ACCENT_3)
        add_text(slide, 8.02, top_y+2.48, 4.4, 1.38, insight[:4] or ['这段代码支撑了当前页的判断。'], size=13.2)
    else:
        add_panel(slide, 7.82, top_y, 4.8, 2.22, '代码理解支撑', title_color=ACCENT_3)
        add_text(slide, 8.02, top_y+0.38, 4.4, 1.72, insight[:5] or ['这页主要是概览，不强行上代码块。'], size=13.2)
        add_panel(slide, 7.82, top_y+2.42, 4.8, 1.53, '源码依据', title_color=ACCENT)
        add_text(slide, 8.02, top_y+2.78, 4.35, 1.08, sources[:4] or ['本页无单独源码列表'], size=11.6)

    add_panel(slide, 0.62, 6.55, 12.02, 0.68, '希望听众带走什么', title_color=WARN, fill=PANEL_3)
    add_text(slide, 0.84, 6.83, 11.58, 0.18, takeaway or ['本页的重点是建立正确心智模型。'], size=14.2, bullet=False)
    footer_right = ' / '.join(sources[:4])
    add_footer(slide, 'Claude Code 源码调研', footer_right)


def function_slide(prs: Presentation, sd: SlideData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, sd.title, '功能 / 实现原理 / 使用技巧 / 代码理解')
    question = parse_bullets(sd.fields.get('这一页要回答的问题', ''))
    if question:
        add_panel(slide, 0.62, 1.5, 12.02, 0.78, '这一页要回答的问题', title_color=ACCENT)
        add_text(slide, 0.82, 1.85, 11.6, 0.25, question, size=14.5, bullet=False)
        top_y = 2.48
    else:
        top_y = 1.55
    add_panel(slide, 0.62, top_y, 3.92, 1.35, '功能', title_color=ACCENT)
    add_text(slide, 0.82, top_y+0.37, 3.5, 0.88, parse_bullets(sd.fields.get('功能', '')) or [strip_links(sd.fields.get('功能', ''))], size=13.6)
    add_panel(slide, 0.62, top_y+1.52, 3.92, 2.18, '实现原理', title_color=ACCENT_2)
    add_text(slide, 0.82, top_y+1.89, 3.5, 1.7, parse_bullets(sd.fields.get('实现原理', '')) or [strip_links(sd.fields.get('实现原理', ''))], size=13.0)
    add_panel(slide, 4.72, top_y, 3.1, 3.7, '使用技巧', title_color=ACCENT_3)
    add_text(slide, 4.92, top_y+0.37, 2.7, 3.18, parse_bullets(sd.fields.get('使用技巧', '')) or [strip_links(sd.fields.get('使用技巧', ''))], size=12.8)
    code = parse_code(sd.fields.get('关键代码片段', ''))
    insight = parse_bullets(sd.fields.get('代码理解支撑', ''))
    sources = markdown_label_list(sd.fields.get('代码支撑', ''))
    if code:
        add_code_box(slide, 8.02, top_y, 4.62, 2.1, code, '关键源码片段')
        add_panel(slide, 8.02, top_y+2.28, 4.62, 1.42, '这段代码说明什么', title_color=WARN)
        add_text(slide, 8.2, top_y+2.62, 4.22, 0.94, insight[:4], size=12.8)
    else:
        add_panel(slide, 8.02, top_y, 4.62, 2.0, '代码理解支撑', title_color=WARN)
        add_text(slide, 8.2, top_y+0.38, 4.22, 1.46, insight[:5], size=12.8)
        add_panel(slide, 8.02, top_y+2.18, 4.62, 1.52, '源码依据', title_color=ACCENT)
        add_text(slide, 8.2, top_y+2.54, 4.22, 1.0, sources[:4] or ['本页无单独源码列表'], size=11.4)
    takeaway = parse_bullets(sd.fields.get('希望听众带走什么', ''))
    add_panel(slide, 0.62, 6.55, 12.02, 0.68, '希望听众带走什么', title_color=WARN, fill=PANEL_3)
    add_text(slide, 0.84, 6.83, 11.58, 0.18, takeaway or ['理解机制之后，才能推出更稳定的使用方式。'], size=14.0, bullet=False)
    add_footer(slide, 'Claude Code 源码调研', ' / '.join(sources[:4]))


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides = parse_outline()
    for sd in slides:
        if sd.num == 1:
            add_cover_slide(prs, sd)
        elif sd.num == 4:
            add_overview_diagram(prs, sd)
        elif sd.num == 10:
            add_workflow_diagram(prs, sd)
        elif sd.num == 17:
            add_capability_slide(prs, sd)
        elif '功能' in sd.fields:
            function_slide(prs, sd)
        else:
            generic_slide(prs, sd)

    prs.save(OUT)
    print(f'Wrote {OUT}')


if __name__ == '__main__':
    build_deck()
