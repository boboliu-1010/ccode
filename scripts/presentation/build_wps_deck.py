from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


TITLE_COLOR = RGBColor(11, 31, 68)
TEXT_COLOR = RGBColor(31, 41, 55)
MUTED_COLOR = RGBColor(71, 85, 105)
ACCENT = RGBColor(32, 82, 149)
ACCENT_2 = RGBColor(221, 234, 251)
BG = RGBColor(248, 250, 252)
LEAD_BG = RGBColor(16, 28, 63)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(219, 228, 240)

FONT = "Arial"


def set_run_font(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_text_frame_style(tf):
    tf.word_wrap = True
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)


def add_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, text, lead=False):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.42), Inches(11.8), Inches(0.75))
    tf = box.text_frame
    set_text_frame_style(tf)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_run_font(r, 26 if not lead else 28, WHITE if lead else TITLE_COLOR, True)
    if not lead:
        line = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.68), Inches(1.18), Inches(2.2), Inches(0.03)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = LINE
        line.line.fill.background()


def add_subtitle(slide, text, lead=False):
    box = slide.shapes.add_textbox(Inches(0.72), Inches(1.4), Inches(11.2), Inches(0.7))
    tf = box.text_frame
    set_text_frame_style(tf)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_run_font(r, 15, WHITE if lead else MUTED_COLOR)


def add_bullets(slide, bullets, top=1.7, left=0.85, width=11.2, height=4.6, color=TEXT_COLOR, size=19):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    set_text_frame_style(tf)
    tf.clear()
    first = True
    for item in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.bullet = True
        p.level = 0
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = item
        set_run_font(r, size, color)


def add_quote(slide, text):
    box = slide.shapes.add_textbox(Inches(0.92), Inches(6.05), Inches(11.0), Inches(0.55))
    tf = box.text_frame
    set_text_frame_style(tf)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_run_font(r, 14, MUTED_COLOR)


def add_box(slide, x, y, w, h, text, fill=WHITE, line=ACCENT, text_color=TITLE_COLOR, size=16, bold=True):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    tf = shape.text_frame
    set_text_frame_style(tf)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_run_font(r, size, text_color, bold)
    return shape


def add_connector(slide, x1, y1, x2, y2):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = ACCENT
    line.line.width = Pt(2)
    try:
        line.line.end_arrowhead = True
    except Exception:
        pass
    return line


def add_lead_slide(prs, title, subtitle, bullets=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LEAD_BG)
    add_title(slide, title, lead=True)
    add_subtitle(slide, subtitle, lead=True)
    if bullets:
        add_bullets(slide, bullets, top=2.3, left=0.9, width=11.2, height=2.2, color=WHITE, size=18)
    return slide


def add_content_slide(prs, title, bullets, quote=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, title)
    add_bullets(slide, bullets, top=1.65, left=0.85, width=11.2, height=4.8)
    if quote:
        add_quote(slide, quote)
    return slide


def add_architecture_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "架构总览")

    add_box(slide, 0.8, 1.7, 2.0, 0.7, "main.tsx", fill=ACCENT_2)
    add_box(slide, 3.3, 1.7, 2.1, 0.7, "Settings / Auth / Policy")
    add_box(slide, 5.9, 1.7, 1.5, 0.7, "Commands")
    add_box(slide, 7.8, 1.7, 1.3, 0.7, "Tools")
    add_box(slide, 9.5, 1.7, 2.0, 0.7, "REPL / QueryEngine")

    add_box(slide, 2.0, 3.4, 2.0, 0.7, "query.ts", fill=ACCENT_2)
    add_box(slide, 4.6, 3.4, 2.2, 0.7, "API Layer")
    add_box(slide, 7.4, 3.4, 2.3, 0.7, "Tool Execution")

    add_box(slide, 3.0, 5.1, 2.0, 0.7, "MCP")
    add_box(slide, 5.5, 5.1, 2.0, 0.7, "Plugins")
    add_box(slide, 8.0, 5.1, 2.0, 0.7, "Skills")

    add_connector(slide, 1.8, 2.4, 3.0, 3.4)
    add_connector(slide, 10.5, 2.4, 9.0, 3.4)
    add_connector(slide, 4.0, 3.75, 4.6, 3.75)
    add_connector(slide, 6.8, 3.75, 7.4, 3.75)
    add_connector(slide, 8.4, 2.4, 8.4, 3.4)
    add_connector(slide, 4.0, 4.1, 4.0, 5.1)
    add_connector(slide, 6.5, 4.1, 6.5, 5.1)
    add_connector(slide, 9.0, 4.1, 9.0, 5.1)

    add_quote(slide, "结论：main.tsx 负责装配，query.ts 负责推进，其余模块提供能力、状态和约束。")


def add_main_flow_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "主执行链路")

    xs = [0.7, 2.5, 4.5, 6.5, 8.7, 10.5]
    labels = ["用户输入", "main.tsx", "QueryEngine", "query.ts", "tool exec", "继续 / 输出"]
    widths = [1.4, 1.5, 1.6, 1.4, 1.5, 1.6]
    for x, w, label in zip(xs, widths, labels):
        add_box(slide, x, 3.0, w, 0.8, label, fill=ACCENT_2 if label in {"query.ts", "main.tsx"} else WHITE)
    for i in range(len(xs) - 1):
        add_connector(slide, xs[i] + widths[i], 3.4, xs[i + 1], 3.4)
    add_quote(slide, "这是一个“模型调用”和“工具调用”反复往返的闭环。")


def add_turn_loop_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "turn loop：核心结构")

    add_box(slide, 0.8, 2.0, 2.0, 0.8, "messages")
    add_box(slide, 3.3, 2.0, 2.3, 0.8, "context shaping")
    add_box(slide, 6.0, 2.0, 2.0, 0.8, "API request")
    add_box(slide, 8.5, 2.0, 2.0, 0.8, "assistant / tool_use")
    add_box(slide, 5.9, 4.2, 2.4, 0.8, "tool_result 回灌", fill=ACCENT_2)

    add_connector(slide, 2.8, 2.4, 3.3, 2.4)
    add_connector(slide, 5.6, 2.4, 6.0, 2.4)
    add_connector(slide, 8.0, 2.4, 8.5, 2.4)
    add_connector(slide, 9.5, 2.8, 7.1, 4.2)
    add_connector(slide, 5.9, 4.6, 1.8, 2.8)

    add_bullets(
        slide,
        [
            "先做 tool result budget、snip、microcompact、collapse、autocompact",
            "再发请求并接收 streaming 输出",
            "如果出现 tool_use，就执行工具并把 tool_result 回写",
            "因此它更像带 recovery 分支的 state machine，而不是线性 pipeline",
        ],
        top=5.1,
        left=0.9,
        width=11.2,
        height=1.4,
        size=16,
    )


def add_compact_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "compact：分层整理，不是一刀切摘要")

    add_box(slide, 0.8, 2.0, 1.5, 0.7, "snip")
    add_box(slide, 2.7, 2.0, 1.8, 0.7, "microcompact")
    add_box(slide, 4.9, 2.0, 1.9, 0.7, "collapse")
    add_box(slide, 7.2, 2.0, 1.8, 0.7, "autocompact")
    add_box(slide, 9.4, 2.0, 2.2, 0.7, "reactive compact")
    for x1, x2 in [(2.3, 2.7), (4.5, 4.9), (6.8, 7.2), (9.0, 9.4)]:
        add_connector(slide, x1, 2.35, x2, 2.35)

    add_bullets(
        slide,
        [
            "要按两条轴来理解：主动整理 vs 被动恢复；轻量缩减 vs full compact",
            "full compact 后保留的不是只有 summary，还包括 boundary、文件恢复、plan_mode、invoked_skills、deferred delta",
            "本质上是在重建可继续工作的最小 context",
        ],
        top=3.6,
        left=0.9,
        width=11.1,
        height=2.0,
        size=17,
    )


def add_context_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "context 不是只靠 messages")

    add_two_col_text(
        slide,
        "Skills / attachments",
        [
            "skill 被扫描到只是 capability 可见",
            "真正调用 skill 时，正文才进入 context",
            "attachments.ts 会按 turn 注入 memories、plan_mode、mailbox、delta 等工作面信息",
        ],
        "关键状态载体",
        [
            "messages：turn-local 工作面",
            "transcript / sidechain：durable history",
            "ToolUseContext：工具执行总线",
            "AppState：宿主共享状态",
        ],
    )


def add_two_col_text(slide, left_title, left_items, right_title, right_items):
    for x, title, items in [(0.8, left_title, left_items), (6.8, right_title, right_items)]:
        hdr = slide.shapes.add_textbox(Inches(x), Inches(1.75), Inches(4.8), Inches(0.5))
        tf = hdr.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        set_run_font(r, 18, ACCENT, True)
        add_bullets(slide, items, top=2.2, left=x, width=4.8, height=3.5, size=17)


def add_tool_pipeline_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "tool call 不是直接执行")

    labels = [
        "schema",
        "validateInput",
        "classifier / hooks",
        "permission",
        "tool.call()",
        "result processing",
    ]
    x = 0.6
    for i, label in enumerate(labels):
        w = 1.8 if i != 2 else 2.1
        add_box(slide, x, 2.7, w, 0.8, label, fill=ACCENT_2 if label == "tool.call()" else WHITE)
        if i < len(labels) - 1:
            add_connector(slide, x + w, 3.1, x + w + 0.2, 3.1)
        x += w + 0.4

    add_bullets(
        slide,
        [
            "真实路径不是 findToolByName -> tool.call 这么短",
            "前面有 schema、hooks、classifier、permission 决策",
            "后面还有 result block、content replacement、post-tool hooks、attachments",
            "autonomy boundary 在 tool pipeline，而不是 UI 表面",
        ],
        top=4.4,
        left=0.9,
        width=11.1,
        height=1.8,
        size=17,
    )


def add_task_runtime_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "task runtime 与 mailbox")

    add_box(slide, 0.9, 2.2, 1.8, 0.8, "User / Main")
    add_box(slide, 3.3, 2.2, 2.0, 0.8, "task runtime")
    add_box(slide, 5.9, 2.2, 2.0, 0.8, "subagent")
    add_box(slide, 8.5, 2.2, 2.1, 0.8, "mailbox / queue")
    add_box(slide, 5.8, 4.4, 2.3, 0.8, "task-notification", fill=ACCENT_2)
    add_connector(slide, 2.7, 2.6, 3.3, 2.6)
    add_connector(slide, 5.3, 2.6, 5.9, 2.6)
    add_connector(slide, 7.9, 2.6, 8.5, 2.6)
    add_connector(slide, 8.5, 3.0, 6.9, 4.4)
    add_connector(slide, 5.8, 4.8, 1.9, 3.0)

    add_bullets(
        slide,
        [
            "task/framework.ts 负责注册、轮询、GC、notification",
            "runAgent.ts 负责 subagent、sidechain transcript、metadata",
            "mailbox / pendingMessages 的投递发生在 turn 边界",
            "这牺牲了一点即时性，但换来了 turn 内状态一致性",
        ],
        top=5.3,
        left=0.9,
        width=11.1,
        height=1.2,
        size=16,
    )


def add_control_planes_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "为什么它更像产品化 runtime")
    add_two_col_text(
        slide,
        "控制面",
        [
            "认证：OAuth / API key / third-party provider",
            "策略：policy_limits、remote managed settings、forceLoginOrgUUID",
            "结论：服务端负责裁决，客户端负责执行",
        ],
        "扩展面与观测",
        [
            "API 层是多 provider 统一适配层",
            "prompt stack 是一层 control plane",
            "MCP / plugins / commands / skills 都是一级扩展面",
            "queryProfiler / analyzeContext / cacheBreakDetection 负责解释系统为何变慢或变贵",
        ],
    )


def add_kernel_vs_host_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "不是一个内核，而是两层")
    add_two_col_text(
        slide,
        "query.ts = runtime kernel",
        [
            "负责一轮怎么跑",
            "推进 model -> tool -> model 闭环",
            "处理 continue / retry / compact / stop",
            "维护合法轨迹而不是纯文本输出",
        ],
        "QueryEngine.ts = conversation host",
        [
            "负责一段会话怎么活",
            "保持多轮 messages 与 usage",
            "预落盘 transcript",
            "做 SDK / headless 输出投影",
        ],
    )


def add_recovery_graph_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "turn loop 更像 recovery graph")

    add_box(slide, 0.8, 2.4, 1.6, 0.75, "messages")
    add_box(slide, 2.9, 2.4, 1.8, 0.75, "sample")
    add_box(slide, 5.2, 2.4, 1.8, 0.75, "tool_use")
    add_box(slide, 7.5, 2.4, 1.8, 0.75, "tool_result")
    add_box(slide, 9.8, 2.4, 1.8, 0.75, "next turn", fill=ACCENT_2)
    for x1, x2 in [(2.4, 2.9), (4.7, 5.2), (7.0, 7.5), (9.3, 9.8)]:
        add_connector(slide, x1, 2.78, x2, 2.78)

    add_box(slide, 3.0, 4.45, 2.2, 0.7, "prompt_too_long")
    add_box(slide, 5.6, 4.45, 2.2, 0.7, "max_output_tokens")
    add_box(slide, 8.2, 4.45, 2.2, 0.7, "reactive compact", fill=ACCENT_2)
    add_connector(slide, 3.8, 3.15, 4.0, 4.45)
    add_connector(slide, 6.1, 3.15, 6.7, 4.45)
    add_connector(slide, 8.9, 5.15, 10.2, 3.15)

    add_bullets(
        slide,
        [
            "它维护的是可继续、可恢复、可压缩、可回放的合法轨迹",
            "关键不是“发一次请求”，而是不断修正同一条 trajectory",
        ],
        top=5.45,
        left=0.9,
        width=11.0,
        height=1.0,
        size=16,
    )


def add_transcript_recovery_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "transcript / recovery 是 durable runtime 骨架")
    add_bullets(
        slide,
        [
            "sessionStorage.ts 保存的不是普通聊天记录，而是带 parentUuid 的消息链",
            "progress message 不进入主链，compact boundary 会进入恢复协议",
            "conversationRecovery.ts 做的不是反序列化，而是把历史修回 API 可继续状态",
            "synthetic continuation 说明系统在恢复“未完成的 runtime”，不是回放聊天 UI",
        ],
        top=1.8,
        left=0.85,
        width=11.2,
        height=3.3,
    )
    add_quote(slide, "一句话：transcript 提供可追溯历史，recovery 把它修回可继续运行的当前状态。")


def add_tool_result_budget_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "tool result budget / content replacement 很关键")
    add_bullets(
        slide,
        [
            "它解决的不是“截断输出”，而是“冻结输出命运”",
            "某个 tool_use_id 一旦决定替换，后面必须始终复用同一 replacement string",
            "大输出先外化到磁盘，再用稳定 preview 进入 context",
            "恢复后还要重放相同 replacement，才能保住 prompt cache prefix",
        ],
        top=1.8,
        left=0.85,
        width=11.2,
        height=3.4,
    )
    add_quote(slide, "这层是 Claude Code 和很多 demo agent 最大的差别之一。")


def add_compact_meaning_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "compact 的价值不是“做摘要”")
    add_two_col_text(
        slide,
        "触发层次",
        [
            "microcompact",
            "snip",
            "context collapse",
            "autocompact",
            "reactive compact",
        ],
        "compact 后保留什么",
        [
            "boundary + summary",
            "messagesToKeep",
            "文件恢复 attachment",
            "plan_mode / invoked_skills / deferred delta",
        ],
    )
    add_quote(slide, "结论：compact 更像“重建可继续工作的最小 context”，不是只留一段总结。")


def add_context_router_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "attachments.ts 是 context router")
    add_bullets(
        slide,
        [
            "每轮真正送给模型的 context，不只是 messages + system prompt",
            "attachments.ts 会按 turn 注入 memories、dynamic skills、plan_mode、deferred_tools_delta、teammate_mailbox 等",
            "很多运行时信息不是常驻消息，而是按 turn 临时重建工作面",
            "这就是为什么它更像“第二调度层”，而不是附件工具箱",
        ],
        top=1.8,
        left=0.85,
        width=11.2,
        height=3.5,
    )


def add_permission_pipeline_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "真正的 autonomy boundary 在 tool pipeline")

    labels = ["schema", "hooks", "classifier", "permission", "tool.call", "post-hooks"]
    x = 0.55
    widths = [1.4, 1.4, 1.6, 1.6, 1.5, 1.6]
    for label, w in zip(labels, widths):
        add_box(slide, x, 2.45, w, 0.78, label, fill=ACCENT_2 if label == "tool.call" else WHITE)
        x += w + 0.32
    x = 0.55
    for w in widths[:-1]:
        add_connector(slide, x + w, 2.84, x + w + 0.32, 2.84)
        x += w + 0.32

    add_bullets(
        slide,
        [
            "tool call 不是 findToolByName -> tool.call 这么短",
            "permissions、hooks、classifier 在这里不是附属模块，而是同一条执行链",
            "这决定了 Claude Code 的 autonomy boundary 不在 UI，而在 runtime pipeline 本身",
        ],
        top=4.15,
        left=0.9,
        width=11.0,
        height=1.7,
        size=16,
    )


def add_task_runtime_deeper_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "tasks / subagents / mailbox 不是附属功能")
    add_bullets(
        slide,
        [
            "Task.ts 先把执行体建模成一等对象：id、type、status、outputFile、outputOffset、notified",
            "task/framework.ts 管注册、轮询、offset、GC、notification",
            "runAgent.ts 会写 sidechain transcript 和 agent metadata，说明 subagent 是独立执行体",
            "mailbox / pendingMessages 的投递发生在 turn 边界，协作语义更像 actor mailbox",
        ],
        top=1.8,
        left=0.85,
        width=11.2,
        height=3.6,
    )


def add_state_carriers_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "文档里最重要的一组判断：状态载体分层")
    add_two_col_text(
        slide,
        "turn-local / host-wide",
        [
            "messages：turn-local 工作面",
            "ToolUseContext：工具执行总线",
            "AppState：宿主共享状态",
        ],
        "durable / on-demand",
        [
            "transcript / sidechain：durable history",
            "attachments / sidecar artifacts：按需重注入",
            "不要把系统简化成“只有消息数组”",
        ],
    )


def add_debugging_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "文档里的另一个高价值产出：调试入口")
    add_two_col_text(
        slide,
        "遇到问题先看哪里",
        [
            "turn loop：query.ts / QueryEngine.ts",
            "context 过大：compact.ts / microCompact.ts / analyzeContext.ts",
            "工具被拒绝：toolExecution.ts / permissions / hooks",
        ],
        "排障价值",
        [
            "它把“读源码”变成“可直接排障”",
            "能快速定位：为什么继续、为什么 compact、为什么拒绝、为什么变慢",
            "这也是走读文档区别于普通目录介绍的地方",
        ],
    )


def add_code_quality_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "代码质量判断也需要进 PPT")
    add_two_col_text(
        slide,
        "优点",
        [
            "工程成熟度高，很多细节来自线上问题",
            "auth、settings、tools、mcp、query 分层总体清楚",
            "对 shell 安全、缓存、401、多进程锁处理扎实",
        ],
        "主要结构债",
        [
            "query.ts 接近 God Loop",
            "ToolUseContext 过胖",
            "cache invariants 分散",
            "permissions / tasks / remote execution 已经上升成 control plane，但 owner 仍分散",
        ],
    )


def add_takeaways_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, "最后只保留最重要的结论")
    add_two_col_text(
        slide,
        "读源码时先抓",
        [
            "main.tsx -> query.ts -> tools",
            "settings -> auth -> api client",
            "policy / managed settings -> 本地执行",
        ],
        "从源码提炼的使用建议",
        [
            "先读相关代码再改",
            "优先复用现有实现",
            "做最小必要改动",
            "验证后如实汇报",
        ],
    )
    add_quote(slide, "一句话：Claude Code 本质上是一个围绕 turn loop 构建的、可被企业管理的终端 agent runtime。")


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_lead_slide(
        prs,
        "Claude Code Source 代码走读",
        "基于源码的 runtime 架构、主调用链与关键设计",
        ["重点：turn loop、compact、tool execution、task runtime", "目标：先建立整体模型，再进入关键代码"],
    )
    add_content_slide(
        prs,
        "分享目录",
        [
            "这份仓库本质上是什么",
            "架构总览与主执行链路",
            "turn loop、compact 与 context",
            "tool pipeline、task runtime 与 mailbox",
            "认证、策略、扩展系统与可观测性",
            "最后总结与阅读建议",
        ],
    )
    add_content_slide(
        prs,
        "这份仓库本质上是什么",
        [
            "不是简单聊天 CLI，也不是薄薄的 SDK demo",
            "更接近一个完整的本地 agent runtime",
            "核心包括：CLI/TUI、会话状态、turn loop、认证/API、策略控制、扩展系统",
            "阅读时不要只把它理解成“调用模型 API 的壳”",
        ],
    )
    add_architecture_slide(prs)
    add_main_flow_slide(prs)
    add_kernel_vs_host_slide(prs)
    add_turn_loop_slide(prs)
    add_recovery_graph_slide(prs)
    add_transcript_recovery_slide(prs)
    add_tool_result_budget_slide(prs)
    add_compact_slide(prs)
    add_compact_meaning_slide(prs)
    add_context_router_slide(prs)
    add_permission_pipeline_slide(prs)
    add_task_runtime_slide(prs)
    add_state_carriers_slide(prs)
    add_debugging_slide(prs)
    add_control_planes_slide(prs)
    add_code_quality_slide(prs)
    add_takeaways_slide(prs)
    add_lead_slide(prs, "谢谢", "这份源码最值得看的，不只是功能，而是它已经为长时间工作的 runtime 付过哪些工程账。")

    out = "docs/zh/00-代码总览与运行时走读-WPS兼容版.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
