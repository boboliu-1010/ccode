from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent.parent
OUT = ROOT / "docs/zh/10-Claude Code 源码深度解读.pptx"

FONT = "Arial"
MONO = "Courier New"

BG = RGBColor(3, 8, 23)
BG_2 = RGBColor(7, 16, 35)
PANEL = RGBColor(10, 21, 43)
PANEL_2 = RGBColor(12, 27, 56)
CODE_BG = RGBColor(4, 12, 28)
TITLE = RGBColor(242, 247, 255)
TEXT = RGBColor(220, 228, 242)
MUTED = RGBColor(144, 162, 188)
ACCENT = RGBColor(56, 189, 248)
ACCENT_2 = RGBColor(167, 139, 250)
ACCENT_3 = RGBColor(45, 212, 191)
LINE = RGBColor(32, 52, 90)
CODE_TEXT = RGBColor(187, 247, 255)
WARN = RGBColor(251, 191, 36)


def set_run_font(run, size, color, bold=False, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def style_tf(tf, top=4, bottom=4, left=8, right=8):
    tf.word_wrap = True
    tf.margin_top = Pt(top)
    tf.margin_bottom = Pt(bottom)
    tf.margin_left = Pt(left)
    tf.margin_right = Pt(right)
    tf.vertical_anchor = MSO_ANCHOR.TOP


def add_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_glow_bar(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.58), Inches(0.34), Inches(12.0), Inches(0.72))
    tf = box.text_frame
    style_tf(tf, 1, 1, 2, 2)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run_font(r, 26, TITLE, True)
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.62), Inches(1.03), Inches(11.8), Inches(0.38))
        tf = sub.text_frame
        style_tf(tf, 1, 1, 2, 2)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = subtitle
        set_run_font(r, 13, MUTED)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.62), Inches(1.39), Inches(2.45), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def add_chip(slide, x, y, w, h, text, fill=PANEL_2, color=ACCENT):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = color
    tf = shape.text_frame
    style_tf(tf, 1, 1, 4, 4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_run_font(r, 11, color, True)
    return shape


def add_panel(slide, x, y, w, h, title, title_color=ACCENT, body_fill=PANEL):
    body = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    body.fill.solid()
    body.fill.fore_color.rgb = body_fill
    body.line.color.rgb = LINE
    hdr = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34))
    hdr.fill.solid()
    hdr.fill.fore_color.rgb = PANEL_2
    hdr.line.fill.background()
    tf = hdr.text_frame
    style_tf(tf, 1, 1, 5, 5)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run_font(r, 11, title_color, True)
    return body


def add_text_block(slide, x, y, w, h, lines, size=17, color=TEXT, bullet=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    style_tf(tf)
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.bullet = bullet
        p.level = 0
        p.space_after = Pt(7)
        r = p.add_run()
        r.text = line
        set_run_font(r, size, color)
    return box


def add_code_box(slide, x, y, w, h, label, code):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = CODE_BG
    box.line.color.rgb = ACCENT
    hdr = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.28))
    hdr.fill.solid()
    hdr.fill.fore_color.rgb = PANEL_2
    hdr.line.fill.background()
    tf = hdr.text_frame
    style_tf(tf, 1, 1, 6, 6)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = label
    set_run_font(r, 10, ACCENT, True, MONO)
    tf = box.text_frame
    style_tf(tf, 15, 6, 10, 10)
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = code
    set_run_font(r, 11.5, CODE_TEXT, False, MONO)
    return box


def add_source_box(slide, x, y, w, h, title, items, color=ACCENT_3):
    add_panel(slide, x, y, w, h, title, title_color=color)
    add_text_block(slide, x + 0.08, y + 0.38, w - 0.16, h - 0.46, items, size=12.5, color=TEXT, bullet=False)


def add_thesis_box(slide, text):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.62), Inches(1.62), Inches(12.02), Inches(0.72))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(7, 22, 46)
    box.line.color.rgb = ACCENT
    tf = box.text_frame
    style_tf(tf, 4, 4, 10, 10)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_run_font(r, 18, TITLE, True)
    return box


def add_footer(slide, left, right=None):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(7.0), Inches(12.0), Inches(0.25))
    tf = box.text_frame
    style_tf(tf, 0, 0, 0, 0)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = left
    set_run_font(r, 10.5, MUTED)
    if right:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = right
        set_run_font(r, 10.5, MUTED)


def connect(slide, x1, y1, x2, y2, color=ACCENT):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(1.6)
    line.line.end_arrowhead = True
    return line


def add_rich_slide(slide, title, subtitle, thesis, bullets, sources, insight, code_label=None, code=None):
    add_bg(slide)
    add_title(slide, title, subtitle)
    add_thesis_box(slide, thesis)
    add_panel(slide, 0.62, 2.58, 7.0, 3.98, "核心内容")
    add_text_block(slide, 0.78, 2.96, 6.7, 3.56, bullets, size=16.5)
    add_source_box(slide, 7.82, 2.58, 4.8, 1.38, "代码支撑", sources)
    add_source_box(slide, 7.82, 5.32, 4.8, 1.24, "代码理解支撑", [insight], color=ACCENT_2)
    if code:
        add_code_box(slide, 7.82, 4.08, 4.8, 1.08, code_label or "源码代码", code)
    add_footer(slide, "Claude Code 源码深度解读")


def add_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG_2)
    add_glow_bar(slide, 0, 0, 13.333, 0.06, ACCENT)
    add_glow_bar(slide, 0, 7.44, 13.333, 0.06, ACCENT_2)
    add_title(
        slide,
        "Claude Code：更好地使用它的源码解读",
        "总体架构、工作流程、功能设计、使用技巧与协作建议",
    )
    add_thesis_box(slide, "这场分享的目标不是完整解释源码，而是利用源码理解，推出更有效的使用方式。")
    add_panel(slide, 0.72, 2.62, 5.68, 2.9, "这场分享关注什么")
    add_text_block(
        slide,
        0.92,
        3.0,
        5.25,
        2.45,
        [
            "先建立 Claude Code 的最小必要架构模型。",
            "再解释一次请求为什么会稳定、漂移或中断。",
            "最后把结论落到产出、使用技巧和协作方式。",
        ],
        size=17,
    )
    add_panel(slide, 6.72, 2.62, 5.9, 2.9, "这场分享希望解决什么")
    add_text_block(
        slide,
        6.94,
        3.0,
        5.42,
        2.45,
        [
            "为什么同样是 Claude Code，有的人越用越顺，有的人越用越乱。",
            "哪些任务最适合它，哪些任务要谨慎交给它。",
            "怎样组织任务，才能让它稳定地产出工程结果。",
        ],
        size=17,
    )
    add_source_box(
        slide,
        0.72,
        5.8,
        11.9,
        0.84,
        "核心判断",
        [
            "Claude Code 不是自由聊天助手，而是一套 terminal agent runtime（终端代理运行时）；理解这一点，使用方式才会自然变稳。"
        ],
        color=WARN,
    )
    add_footer(slide, "Claude Code 源码深度解读")


def add_architecture_diagram_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "总体架构：Claude Code 的主执行链", "先建立最小必要模型，再进入工作流程与使用技巧。")
    add_thesis_box(slide, "用户输入不是直接发给模型，而是先穿过会话宿主、轮次循环、工具执行链和控制面。")
    labels = [
        ("User Prompt", 0.8, 2.7, 1.8, 0.68, ACCENT),
        ("Prompt Stack\n提示词栈", 2.9, 2.7, 2.0, 0.82, ACCENT),
        ("QueryEngine\n会话宿主", 5.25, 2.64, 2.0, 0.9, ACCENT_2),
        ("query.ts\nTurn Loop", 7.65, 2.64, 2.0, 0.9, ACCENT_2),
        ("Tool Pipeline\n工具执行流水线", 10.02, 2.58, 2.32, 1.02, ACCENT_3),
    ]
    for text, x, y, w, h, color in labels:
        shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = PANEL
        shp.line.color.rgb = color
        tf = shp.text_frame
        style_tf(tf, 4, 4, 8, 8)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        for idx, line in enumerate(text.split("\n")):
            r = p.add_run() if idx == 0 else p.add_run()
            r.text = ("" if idx == 0 else "\n") + line
            set_run_font(r, 16 if idx == 0 else 13, TITLE if idx == 0 else color, idx == 0)
    connect(slide, 2.6, 3.04, 2.9, 3.04)
    connect(slide, 4.9, 3.04, 5.25, 3.04)
    connect(slide, 7.25, 3.04, 7.65, 3.04)
    connect(slide, 9.65, 3.04, 10.02, 3.04)

    env = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.9), Inches(4.45), Inches(3.35), Inches(1.0))
    env.fill.solid()
    env.fill.fore_color.rgb = PANEL
    env.line.color.rgb = ACCENT_3
    tf = env.text_frame
    style_tf(tf, 5, 5, 8, 8)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Tools / Tasks / MCP\n执行环境与扩展能力"
    set_run_font(r, 15, TITLE, True)
    connect(slide, 11.18, 3.62, 10.55, 4.45, ACCENT_3)

    cp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.84), Inches(4.38), Inches(3.55), Inches(1.08))
    cp.fill.solid()
    cp.fill.fore_color.rgb = PANEL
    cp.line.color.rgb = WARN
    tf = cp.text_frame
    style_tf(tf, 5, 5, 8, 8)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Settings / Auth / Policy / Prompt\n控制面"
    set_run_font(r, 15, TITLE, True)

    obs = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(4.78), Inches(4.38), Inches(3.52), Inches(1.08))
    obs.fill.solid()
    obs.fill.fore_color.rgb = PANEL
    obs.line.color.rgb = ACCENT
    tf = obs.text_frame
    style_tf(tf, 5, 5, 8, 8)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Transcript / Recovery / Observability\n续航与观测"
    set_run_font(r, 15, TITLE, True)
    connect(slide, 5.78, 4.38, 5.78, 3.55, WARN)
    connect(slide, 8.3, 4.38, 8.3, 3.55, ACCENT)
    add_footer(slide, "架构图回答：用户输入为什么不会直接等于模型输出。")


def add_workflow_diagram_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "工作流程：一次请求如何流过 Claude Code", "流程图的重点不是步骤数量，而是它为什么能够持续推进。")
    add_thesis_box(slide, "Claude Code 的强项不是一次回答，而是把请求变成一个可推进、可约束、可恢复的工作过程。")
    steps = [
        ("1", "输入进入\nCLI / REPL", 0.88, 2.65),
        ("2", "组装 settings / auth /\ntools / app state", 3.05, 2.65),
        ("3", "进入 QueryEngine\n构造会话状态", 5.45, 2.65),
        ("4", "query.ts\n推进 turn loop", 7.9, 2.65),
        ("5", "Tool Pipeline\n执行工具链", 10.15, 2.65),
        ("6", "结果回灌 /\n继续 / compact / recovery", 5.45, 4.55),
    ]
    for num, text, x, y in steps:
        chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(0.42), Inches(0.42))
        chip.fill.solid()
        chip.fill.fore_color.rgb = ACCENT
        chip.line.fill.background()
        tf = chip.text_frame
        style_tf(tf, 1, 1, 1, 1)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = num
        set_run_font(r, 13, BG, True)
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x + 0.5), Inches(y - 0.02), Inches(1.55 if num != "6" else 2.1), Inches(0.82))
        box.fill.solid()
        box.fill.fore_color.rgb = PANEL
        box.line.color.rgb = ACCENT if num in {"1", "2", "4"} else ACCENT_2
        tf = box.text_frame
        style_tf(tf, 4, 4, 6, 6)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        set_run_font(r, 13.5, TITLE, True)
    connect(slide, 2.92, 3.05, 3.05, 3.05)
    connect(slide, 5.28, 3.05, 5.45, 3.05)
    connect(slide, 7.73, 3.05, 7.9, 3.05)
    connect(slide, 10.0, 3.05, 10.15, 3.05)
    connect(slide, 11.95, 3.45, 8.05, 4.55, ACCENT_3)
    connect(slide, 6.55, 4.55, 6.55, 3.47, ACCENT_3)
    add_source_box(
        slide,
        0.9,
        5.85,
        11.7,
        0.78,
        "代码理解支撑",
        ["这条流程由 main.tsx、QueryEngine.ts、query.ts 和 toolExecution.ts 共同定义；它天然是一条多轮执行链，而不是一次性问答。"],
        color=ACCENT_3,
    )
    add_footer(slide, "工作流程图回答：为什么 Claude Code 更像“推进器”，而不是“聊天器”。")


def add_capability_map_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "功能列表：从用户视角看 Claude Code 的能力面", "能力多本身不重要，关键是它们能围绕主执行链协同。")
    add_thesis_box(slide, "Claude Code 的价值不在单个功能，而在交互、执行、扩展、控制、续航和观测能否围绕同一条主链协同。")
    cards = [
        ("交互", ["CLI / REPL", "commands", "Plan Mode", "structured output"], 0.78, 2.62, ACCENT),
        ("执行", ["Read / Edit / Bash", "tool pipeline", "validation", "result processing"], 3.28, 2.62, ACCENT_2),
        ("扩展", ["skills", "MCP", "plugins", "remote capability"], 5.78, 2.62, ACCENT_3),
        ("控制", ["settings", "auth", "policy", "prompt stack"], 8.28, 2.62, WARN),
        ("续航", ["transcript", "compact", "recovery", "task runtime"], 10.78, 2.62, ACCENT),
        ("可观测性", ["query profiler", "cache break", "analyze context"], 4.98, 5.05, ACCENT_2),
    ]
    for title, lines, x, y, color in cards:
        add_panel(slide, x, y, 2.05, 1.85, title, title_color=color)
        add_text_block(slide, x + 0.1, y + 0.4, 1.85, 1.4, lines, size=12.5)
    add_footer(slide, "功能列表的重点：它最适合产出什么结果，以及哪些能力会直接影响使用方式。")


def add_section_intro(prs, title, subtitle, bullets, tag):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG_2)
    add_title(slide, title, subtitle)
    add_thesis_box(slide, tag)
    add_panel(slide, 0.85, 2.55, 11.6, 3.55, "本部分关注")
    add_text_block(slide, 1.08, 2.95, 11.1, 3.05, bullets, size=20)
    add_footer(slide, "Claude Code 源码深度解读")


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_cover(prs)
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "这场分享真正要解决的问题",
        "为什么同样是 Claude Code，有的人越用越顺，有的人越用越乱。",
        "问题的关键不在“会不会用模型”，而在能不能顺着 Claude Code 的默认运行方式组织任务。",
        [
            "Claude Code 不是自由聊天助手，而是围绕工程任务设计的执行型系统。",
            "输入方式、任务边界、验证要求和风险约束，都会直接影响输出质量。",
            "如果只把它当聊天工具来用，就会频繁遇到范围漂移、验证不足和边界失控。",
            "如果顺着它的工作链条组织任务，Claude Code 会更像一个稳定的工程协作体。",
        ],
        ["docs/zh/09-Claude Code 源码深度解读-文档.md", "docs/zh/10-Claude Code 源码深度解读-分享提纲.md"],
        "这场分享的主次关系是：源码理解做铺垫，重点落在产出、用法和协作方式。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "先给结论：Claude Code 到底是什么",
        "先建立一个足够准确的工程定义，后面所有技巧才有依托。",
        "Claude Code 更像 terminal agent runtime（终端代理运行时），而不是普通聊天 CLI 或薄薄一层 API wrapper。",
        [
            "main.tsx 负责启动装配：配置、认证、插件、技能、MCP、命令和工具。",
            "QueryEngine.ts 负责会话宿主：消息、usage、file state、transcript。",
            "query.ts 负责轮次循环：上下文整理、模型采样、工具回灌、继续 / 重试 / 压缩。",
            "toolExecution.ts 负责工具执行链：parse、validate、hooks、permissions、post-process。",
        ],
        ["src/main.tsx", "src/QueryEngine.ts", "src/query.ts", "src/services/tools/toolExecution.ts"],
        "装配、宿主、执行和控制面同时存在，才说明它更接近 runtime，而不是普通 CLI wrapper。",
    )
    add_architecture_diagram_slide(prs)
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "Claude Code 和 demo agent 的本质差别",
        "理解它为什么像生产系统，决定了后面讲什么才有价值。",
        "demo agent 的中心是“模型调工具”；Claude Code 的中心是“让一条执行轨迹在复杂条件下仍可持续”。",
        [
            "它显式处理恢复、压缩、权限链、多执行体、观测与诊断。",
            "它不是只追求“能跑一次”，而是追求“在长任务里还能继续跑”。",
            "对话、工具、任务、扩展和控制面共同围绕同一条主执行链工作。",
            "真正稀缺的价值，不是功能多少，而是已经为长期工作付过工程账。",
        ],
        ["src/query.ts", "src/utils/conversationRecovery.ts", "src/services/tools/toolExecution.ts", "src/utils/task/framework.ts"],
        "query.ts 的迁移分支、recovery、task framework 和 tool pipeline 一起说明它不属于 demo 形态。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "main.tsx、QueryEngine.ts、query.ts 的分工",
        "这是理解 Claude Code 的三个核心入口。",
        "一个管装配，一个管寿命，一个管推进；把它们混在一起，后面的逻辑就很难讲清。",
        [
            "main.tsx：把 settings、auth、plugins、skills、MCP、commands、tools 组装成可运行系统。",
            "QueryEngine.ts：持有 mutableMessages、usage、readFileState 等会话状态。",
            "query.ts：持有 turn state、toolUseContext、transition 等轮次推进状态。",
            "这种分工决定了 Claude Code 更像“会话型 runtime”，不是一次性问答器。",
        ],
        ["src/main.tsx", "src/QueryEngine.ts", "src/query.ts"],
        "先把装配、宿主、执行三层拆开看，后面谈工作流程和使用技巧才会自然。",
        "源码代码",
        "export class QueryEngine {\n  private mutableMessages: Message[]\n  private totalUsage: NonNullableUsage\n}\n\ntype State = {\n  messages: Message[]\n  turnCount: number\n  transition: Continue | undefined\n}",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "控制面为什么重要",
        "Claude Code 的行为边界，不是只由模型决定。",
        "settings、auth、policy limits、managed settings、prompt stack 共同构成控制面，持续裁决系统行为。",
        [
            "settings 决定功能开关、路径、模型选择和本地行为。",
            "auth 决定身份来源和可调用后端。",
            "policy limits 和 managed settings 决定组织级限制和远程托管行为。",
            "system prompt / dynamic prompt sections 决定当前任务被怎样解释。",
        ],
        ["src/utils/settings/settings.ts", "src/utils/auth.ts", "src/services/policyLimits/index.ts", "src/utils/systemPrompt.ts"],
        "很多“为什么它这样做”的答案不在模型本身，而在控制面如何裁决这个 runtime。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "为什么前面这几页必须先讲",
        "源码理解不是目的，但没有这部分，后面的技巧就会沦为经验帖。",
        "先建立系统分层和控制面模型，后面讲使用方式时，才能解释“为什么这样做更稳”。",
        [
            "总体架构回答：Claude Code 是什么。",
            "主执行链回答：请求如何流过系统。",
            "控制面回答：谁在裁决边界。",
            "这些前提一旦成立，工作流程、产出形态和使用技巧都会顺势成立。",
        ],
        ["src/main.tsx", "src/QueryEngine.ts", "src/query.ts", "src/utils/systemPrompt.ts"],
        "源码理解在这里的作用，是为使用方法提供结构支撑，而不是单独做代码炫技。",
    )

    add_section_intro(
        prs,
        "第二部分：工作流程",
        "理解一次请求如何推进，才能解释什么样的任务组织方式更有效。",
        [
            "一次请求不是“用户说一句，模型回一句”。",
            "它会先过装配和控制面，再进入 QueryEngine 和 turn loop。",
            "主流程中还有工具执行链、压缩、恢复和长任务续航机制。",
        ],
        "工作流程是后面所有使用技巧的因果背景。",
    )
    add_workflow_diagram_slide(prs)
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "一次请求从哪里开始",
        "在模型真正被调用之前，Claude Code 已经做了不少事。",
        "输入进入系统后，先过 settings、auth、policy 和工具装配，再真正进入会话与轮次推进。",
        [
            "CLI / REPL 接收请求后，会先确定环境、配置和身份。",
            "main.tsx 会构造 tools、commands、app state 和扩展能力面。",
            "只有这些前置条件成立，请求才会进入 QueryEngine 和 query.ts。",
            "这解释了为什么同样一句话，在不同环境中可能有不同表现。",
        ],
        ["src/main.tsx", "src/QueryEngine.ts"],
        "请求处理在模型调用前就已经开始，Claude Code 不是单纯的模型转发器。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "turn loop（轮次循环）的基本闭环",
        "Claude Code 的强项不是“回答”，而是“推进”。",
        "一轮 loop 的基本闭环是：当前 messages -> context shaping -> API request -> assistant / tool_use -> tool_result 回灌 -> next turn。",
        [
            "当前 messages 只是起点，不是最终上下文。",
            "系统会先做 context shaping，把 prompt、attachments、memory 和能力面整理出来。",
            "模型返回 assistant message 或 tool_use 后，工具结果会回灌进消息链。",
            "下一轮不是重新开始，而是在前一轮轨迹上继续推进。",
        ],
        ["src/query.ts"],
        "query.ts 围绕 messages、tool context 和 transition 反复推进，因此 loop 是文件真实形态，而不是分析者措辞。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "为什么它不是线性流水线",
        "如果把 Claude Code 当普通工具调用链看，很多设计都会被误读。",
        "Claude Code 更像 recovery graph（恢复图），而不是简单 pipeline。",
        [
            "主流程中途要处理 prompt_too_long、max_output_tokens、continuation、stop hook、reactive compact。",
            "这些不是边角异常处理，而是主流程的一部分。",
            "系统一直在维持一条可持续推进、可恢复、可压缩的执行轨迹。",
            "这也是它比 demo agent 更像生产系统的原因之一。",
        ],
        ["src/query.ts"],
        "continuation、compact 和 recovery 分支在 query.ts 里是正向路径的一部分，而不是附属脚本。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "context shaping（上下文整理）是怎么发生的",
        "模型每一轮看到的，不是静态聊天记录，而是动态构造出来的工作面。",
        "Claude Code 会把 prompt stack、attachments、relevant memories、invoked skills 和 reminders 组合成当前轮上下文。",
        [
            "system prompt 自身就是可分层、可动态拼接的。",
            "attachments 会按当前轮需求注入 memories、skill delta、task messages 和 reminders。",
            "compact 之后，系统还会把必要的工作面信息重新带回来。",
            "这解释了为什么同一会话中，不同轮次的“上下文视图”其实是变化的。",
        ],
        ["src/utils/systemPrompt.ts", "src/utils/attachments.ts"],
        "用户写 prompt，不是在给静态聊天机器人发消息，而是在参与当前轮工作面的构造。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "tool pipeline（工具执行流水线）如何推进",
        "工具调用不是直接 tool.call()，而是一条被审查、被约束、被记录的执行链。",
        "tool pipeline 的重点不是把命令最快打出去，而是确保动作有边界、结果可回灌、状态可继续。",
        [
            "工具输入先过 schema parse 和 validate。",
            "pre hooks、permissions、classifier 会在真正执行前介入。",
            "执行后还有 post process 和结果回写。",
            "因此用户把边界写清楚，会直接帮助这条执行链更稳定工作。",
        ],
        ["src/services/tools/toolExecution.ts"],
        "真正的自主边界位于 runtime 内部，而不是 UI 上一句“是否继续”。",
        "源码代码",
        "const parsedInput = tool.inputSchema.safeParse(input)\n...\nrunPreToolUseHooks(...)\n...\nresolveHookPermissionDecision(...)\n...\ntool.call(...)",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "Transcript / Recovery / Compact 在流程里的位置",
        "长期工作能力不是附属特性，而是这条工作流程的续航机制。",
        "transcript 保存可恢复消息链，recovery 负责回到 API 可继续状态，compact 负责重建还能工作的最小 context。",
        [
            "transcript 不是普通聊天记录，而是带 parentUuid 的因果链。",
            "recovery 会修 unresolved tool use、thinking block 和 continuation 状态。",
            "compact 不只是摘要，而是保留 boundary、tail messages 和 attachments。",
            "这三层一起，才解释了为什么 Claude Code 能做长任务和断点续做。",
        ],
        ["src/utils/sessionStorage.ts", "src/utils/conversationRecovery.ts", "src/services/compact/compact.ts"],
        "长期工作是底层显式目标，而不是幸运地“刚好能跑很久”。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "工作流程小结",
        "理解工作流不是为了背模块，而是为了知道怎样组织任务更符合系统本性。",
        "Claude Code 会把请求变成一个可推进、可约束、可恢复的工作过程；用户越顺着这套过程组织任务，结果越稳定。",
        [
            "一次请求天然是多轮的，不是一句输入对应一句输出。",
            "工具链、压缩、恢复和观测都在主流程里，不是边角功能。",
            "长期工作能力来自持续维护轨迹，而不是持续“记住所有细节”。",
            "这也是后面讲产出和技巧时最重要的背景。",
        ],
        ["src/query.ts", "src/services/tools/toolExecution.ts", "src/utils/conversationRecovery.ts"],
        "工作流程越清楚，越知道什么样的任务和提示方式会更有效。",
    )

    add_section_intro(
        prs,
        "第三部分：功能列表与产出",
        "重点不在功能名本身，而在它们能产出什么工程结果。",
        [
            "从用户视角，Claude Code 的价值应该用“产出形态”来衡量。",
            "哪些任务适合它、哪些任务不适合它，都可以从运行时结构解释。",
            "功能面和产出形态对应起来，后面的使用建议才不会空。",
        ],
        "功能列表只是索引，真正重要的是这些功能如何转化为稳定产出。",
    )
    add_capability_map_slide(prs)
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "从用户视角看，Claude Code 最典型的产出是什么",
        "如果只把它理解成“会回答问题”，就会低估它的真正价值。",
        "Claude Code 最擅长的不是聊天答案，而是工程结果：代码、解释、验证、计划、结构化结论和工作卡。",
        [
            "代码修改：局部修复、局部重构、多文件联动更新。",
            "代码解释与定位：帮你形成“问题在哪、为什么在这里”的结论。",
            "命令执行与验证：把改动、日志、测试和结论串起来。",
            "计划与任务拆解：先探索，再计划，再推进实现。",
            "结构化输出：把最终结果沉淀成更可消费的形式。",
        ],
        ["src/commands.ts", "src/Tool.ts", "src/utils/messages.ts", "src/QueryEngine.ts"],
        "commands、Edit/Read/Bash、Plan Mode 和 structured output 一起说明：它默认就是围绕工程结果设计的。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "哪些任务特别适合 Claude Code",
        "判断一项任务是否适合它，关键看是否贴合它的主能力链。",
        "越像“需要读代码、调工具、形成结果、可分阶段推进”的任务，Claude Code 越容易发挥。",
        [
            "局部修复和局部重构：边界清楚，适合最小改动和验证。",
            "多文件追踪与定位：天然需要读代码、做关联、形成解释。",
            "命令执行与验证：适合把改动和测试结论串在一起。",
            "需要拆阶段推进的任务：适合 Plan Mode、tasks 和 recovery。",
        ],
        ["src/query.ts", "src/services/tools/toolExecution.ts", "src/utils/messages.ts", "src/utils/task/framework.ts"],
        "这些任务之所以适合，是因为它们天然贴合 prompt stack、tool pipeline、task runtime 和 recovery 的优势区间。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "哪些任务不适合用 Claude Code 直接硬做",
        "Claude Code 不是不能做难任务，而是难任务更需要结构化边界。",
        "目标极度模糊、范围完全不受限、输出形式没有约束、需要它自行理解隐含边界的任务，会明显放大不稳定性。",
        [
            "完全模糊的探索：系统缺少稳定工作面。",
            "范围极大的全局变更：compact 和 recovery 代价会快速上升。",
            "高风险但没有确认规则的动作：permission chain 会变得脆弱。",
            "高度依赖隐性业务背景的任务：prompt stack 很难替用户补齐这些边界。",
        ],
        ["src/constants/prompts.ts", "src/services/tools/toolExecution.ts", "src/utils/permissions/permissions.ts"],
        "模糊任务不是“模型太弱”，而是 prompt stack、tool pipeline 和权限链都更容易进入不稳定状态。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "为什么它能稳定地产出这些结果",
        "稳定产出不是模型自行收敛出来的，而是由多条运行时链路共同支持的。",
        "默认角色、工具执行链、续航链和任务链一起存在，决定了 Claude Code 更容易产出“工作结果”，而不是发散回答。",
        [
            "默认角色是工程代理，天然倾向于处理软件工程任务。",
            "工具执行受约束，动作先被审查再被执行。",
            "结果可以进入 transcript / compact / recovery，任务不会轻易断掉。",
            "plan、task、memory 让复杂工作具备多阶段推进能力。",
        ],
        ["src/constants/prompts.ts", "src/services/tools/toolExecution.ts", "src/utils/sessionStorage.ts", "src/utils/messages.ts"],
        "稳定产出来自运行时结构，不只是模型强；这也是为什么用法会直接影响结果质量。",
    )
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "用户视角功能小结",
        "先理解产出形态，再谈技巧，效果会更稳。",
        "Claude Code 最适合推进工程任务，特别适合产出“代码 + 解释 + 验证 + 计划”；它不适合完全模糊、完全无边界的自由发挥。",
        [
            "功能多本身不重要，关键是它们能否协同产出工程结果。",
            "适合的任务天然贴合它的主执行链。",
            "不适合的任务通常意味着边界不清、风险不明或工作面不可控。",
            "下一部分开始逐个讲最重要的功能，以及这些功能对应的使用技巧。",
        ],
        ["src/commands.ts", "src/services/tools/toolExecution.ts", "src/utils/messages.ts"],
        "从这里往后，重点不再是“它有什么功能”，而是“这些功能为什么会导出某种更优的使用方式”。",
    )

    add_section_intro(
        prs,
        "第四部分：重要功能详细介绍",
        "讲实现原理的目的，是推出更有效的使用方式，而不是单独炫耀代码。",
        [
            "每个重点功能都会回答三件事：它是什么、它如何工作、用户应该如何顺着它来用。",
            "所有使用技巧都必须能回到代码机制，而不是停在经验层面。",
            "重点功能讲清以后，后面的流程建议才会足够稳。",
        ],
        "真正有价值的功能，不是功能名，而是它背后的运行时含义。",
    )
    detail_slides = [
        ("Prompt Stack（提示词栈）", "默认角色如何被定义", "system prompt、agent prompt、dynamic sections 共同定义 Claude Code 的默认角色和行为边界。", ["default prompt 定义用户请求首先是 software engineering tasks。", "systemPrompt.ts 负责拼接 override、coordinator、agent 和 dynamic sections。", "prompt stack 决定了 Claude Code 更喜欢工程化表达，而不是闲聊式表达。", "因此任务越像工程任务，越能触发系统默认最佳状态。"], ["src/utils/systemPrompt.ts", "src/constants/prompts.ts"], "默认 system prompt 直接把用户请求定义成工程任务，所以“任务表达越工程化越稳”不是经验，而是系统角色设定的直接结果。", "源码提示词", "The user will primarily request you to perform software engineering tasks."),
        ("QueryEngine（会话宿主）", "Claude Code 为什么像会话型系统", "QueryEngine 持有 mutableMessages、usage、file state 和 transcript 相关状态，因此它像 conversation host，而不是一次性 API wrapper。", ["它负责一整段会话怎么活，而不只是一次模型调用。", "它会在进入 query.ts 前先装配 processUserInput context。", "它也是 SDK / headless 输出的收口层。", "这说明 Claude Code 的核心不是“问一次”，而是“让一段会话持续存在”。"], ["src/QueryEngine.ts"], "持久化消息、usage 聚合和 transcript 预落盘都发生在 QueryEngine 这一层，因此它更像会话宿主。", "源码代码", "export class QueryEngine {\n  private mutableMessages: Message[]\n  private totalUsage: NonNullableUsage\n  private readFileState: FileStateCache\n}"),
        ("Turn Loop（轮次循环）", "Claude Code 为什么擅长推进", "query.ts 保存的是跨轮迁移状态，不只是消息列表，因此它天然适合可继续推进的任务。", ["State 里不仅有 messages，还有 toolUseContext、pending summary、stopHookActive 和 transition。", "每一轮都会重新整理上下文、发起采样、执行工具并继续推进。", "它的强项不是一次性回答，而是围绕执行轨迹持续前进。", "因此大任务要拆阶段，小任务也要给清楚边界。"], ["src/query.ts"], "State 结构说明它维护的是 runtime 状态图，而不是静态消息数组。", "源码代码", "type State = {\n  messages: Message[]\n  toolUseContext: ToolUseContext\n  turnCount: number\n  transition: Continue | undefined\n}"),
        ("Tool Pipeline（工具执行流水线）", "为什么工具调用不是直接执行", "schema parse -> validate -> hooks -> permissions -> call -> post process，这条链定义了 Claude Code 的真实自主边界。", ["执行不是直接 tool.call()，而是先被解析、审查和裁决。", "hooks 和 permissions 不是附属物，而是执行链的一部分。", "tool result 会被标准化并回灌到主轨迹。", "所以用户把边界写清楚，会直接帮助这条执行链更稳定。"], ["src/services/tools/toolExecution.ts"], "真正的“是否代表用户行动”发生在工具执行链内部，而不是在 UI 里口头提醒。", "源码代码", "const parsedInput = tool.inputSchema.safeParse(input)\n...\nrunPreToolUseHooks(...)\n...\nresolveHookPermissionDecision(...)\n...\ntool.call(...)"),
        ("BashTool", "最强能力为什么也最需要边界", "BashTool 提供最强执行能力，但周围专门拆出了 sandbox、只读识别、路径校验和破坏性判断等安全层。", ["Bash 在系统里不是默认首选，而是高能力、高风险通道。", "它会单独处理路径合法性、只读命令和破坏性命令。", "长输出还会触发额外的结果治理。", "因此 Bash 更适合做兜底，而不是默认第一选择。"], ["src/tools/BashTool/BashTool.tsx", "src/tools/BashTool/prompt.ts"], "安全子模块的存在本身就说明 BashTool 被系统视为高风险能力，而不是普通执行路径。"),
        ("Compact（上下文压缩）", "长会话为什么不会直接拖垮系统", "compact 保留的不只是 summary，还有 boundary、tail messages 和 attachments，因此它在重建工作面，而不是简单做摘要。", ["microcompact、autocompact 和 reactive compact 共同存在。", "compact 之后，系统仍然要保住必要的工作面信息。", "用户不能假设模型会永久持有所有历史细节。", "更稳的方式是要求阶段性产出，而不是指望全量记忆。"], ["src/services/compact/compact.ts", "src/services/compact/microCompact.ts"], "compact 的目标是重建可工作 context，而不是把历史压成一段漂亮摘要。"),
        ("Transcript / Recovery（会话记录 / 恢复）", "为什么长任务和断点续做是它的强项", "recovery 会主动修 unresolved tool use、thinking、continuation 等状态，这种深度说明“继续工作”是底层能力。", ["transcript 保存的是可恢复消息链。", "recovery 会主动过滤和修补不合法状态。", "synthetic continuation 等机制直接服务于恢复。", "这使得 Claude Code 适合长任务和断点续做。"], ["src/utils/sessionStorage.ts", "src/utils/conversationRecovery.ts"], "如果 transcript 只是聊天记录，这些修补逻辑没有必要；正因为它是 runtime 状态载体，这些逻辑才成立。", "源码代码", "const filteredToolUses = filterUnresolvedToolUses(migratedMessages)\nconst filteredThinking =\n  filterOrphanedThinkingOnlyMessages(filteredToolUses)"),
        ("Content Replacement（内容替换）", "大工具输出为什么不会直接把上下文拖垮", "系统用 ContentReplacementState 维护 seenIds 和 replacements，说明它在保护“看过的前缀不能漂”。", ["大结果会落盘，只把稳定 preview 留在上下文里。", "同一个 tool_use_id 的 replacement 一旦确定，后面不能随便变。", "resume 后还要重放 replacement。", "因此用户更应该关注结果和结论，而不是要求模型永远记住全部大输出。"], ["src/utils/toolResultStorage.ts", "src/services/compact/microCompact.ts"], "这层设计说明 Claude Code 把 prompt cache prefix stability 当成一等约束。", "源码代码", "export type ContentReplacementState = {\n  seenIds: Set<string>\n  replacements: Map<string, string>\n}"),
        ("Skills（技能）", "为什么 skill 不是一段快捷短语", "skills 会按路径、上下文和 compact 结果激活与保留，因此它们更像 capability artifact，而不是静态 prompt 模板。", ["skills 不是简单的“快捷指令”。", "它们会参与能力面暴露和当前轮上下文构造。", "compact 后 invoked skills 仍会被保留。", "这意味着 skill 更适合有明显任务边界的场景。"], ["src/utils/skills/loadSkillsDir.ts", "src/utils/attachments.ts"], "只要一个能力对象会被按路径和上下文激活，它就已经不再只是文本模板。"),
        ("Attachments（上下文附件）", "为什么 Claude Code 不是只看聊天记录", "attachments 逐轮注入 memories、skill delta、task messages 和 reminders，说明上下文是逐轮构造的工作面。", ["当前轮上下文来自聊天记录、prompt stack 和 attachments 的共同作用。", "attachments 负责把额外但关键的信息带回工作面。", "这也是 Claude Code 能维持复杂任务连续性的原因之一。", "因此用户组织任务的方式，会直接影响模型看到什么。"], ["src/utils/attachments.ts"], "如果把 Claude Code 当成只看消息历史的聊天机器人，就会误解很多行为。"),
        ("Permissions / Hooks / Classifier", "真正的自主边界在哪里", "permissions、hooks 和 classifier 都位于工具执行链中，因此真正的自主边界在 runtime 内部，而不是外部口头约束。", ["allow / ask / deny 规则会在执行链里生效。", "pre/post hooks 可以在真正执行前后插入行为。", "classifier 会参与某些自动模式下的裁决。", "这解释了为什么模糊授权会显著放大不确定性。"], ["src/services/tools/toolExecution.ts", "src/utils/permissions/permissions.ts"], "系统内部对风险动作是单独建模的，因此用户也必须把确认规则显式写出来。"),
        ("Tasks / Subagents / Mailbox", "为什么 Claude Code 适合持续推进而不只是一次性答复", "task registry、agent metadata、mailbox 和 sidechain transcript 同时存在，说明它已经有多执行体 runtime 的雏形。", ["复杂任务可以拆阶段、拆子任务，而不是硬塞成一次答复。", "task runtime 让异步执行体具备生命周期。", "subagent 和 mailbox 让协作语义开始形成。", "因此 Claude Code 更像工作推进器，而不是单轮聊天器。"], ["src/Task.ts", "src/utils/task/framework.ts", "src/tools/AgentTool/runAgent.ts"], "任务系统的存在，直接扩大了它适合承担的任务形态。"),
        ("MCP / Plugins / Remote Capability", "Claude Code 的上限为什么不只取决于模型", "MCP、plugins 和 remote capability 各有接入点，说明扩展能力在系统里是一级设计目标。", ["Claude Code 的能力面可以通过 MCP 和插件持续扩展。", "这意味着“能不能接入合适工具”会直接影响它的上限。", "remote capability 进一步把它从本地代理推向能力枢纽。", "因此使用时要一起看模型能力和可接入能力面。"], ["src/services/mcp/client.ts", "src/utils/plugins/pluginLoader.ts"], "扩展能力不是附属物，而是系统原生设计目标；这解释了为什么 Claude Code 的价值常常来自“模型 + 工具面”。"),
        ("Observability（可观测性）", "为什么它更像生产系统", "query profiler、cache break detection 和 analyzeContext 让系统不只会跑，还能解释自己为什么这样跑。", ["性能、上下文成本和缓存命中都不是黑盒。", "可观测性让团队能诊断慢在哪、贵在哪、坏在哪。", "这对于长任务和复杂工具链尤其重要。", "对使用者来说，意味着系统行为更可解释。"], ["src/utils/queryProfiler.ts", "src/services/api/promptCacheBreakDetection.ts", "src/utils/analyzeContext.ts"], "没有可观测性，就没有可演化的 harness；这也是它和 demo 代码的明显差别。"),
    ]
    for args in detail_slides:
        add_rich_slide(prs.slides.add_slide(prs.slide_layouts[6]), *args)
    add_rich_slide(
        prs.slides.add_slide(prs.slide_layouts[6]),
        "重要功能小结",
        "真正的“使用技巧”，必须回到运行时机制上理解。",
        "Prompt Stack、Turn Loop、Tool Pipeline、Compact、Recovery 和 Tasks 一起决定了系统如何继续工作，也决定了用户应该怎样更稳地使用它。",
        [
            "源码解释了为什么工程化表达更稳、为什么边界要写清、为什么长任务要分阶段。",
            "真正的技巧不是“给它几句神奇提示词”，而是顺着 runtime 的工作方式组织任务。",
            "功能细节讲得越清楚，后面的流程建议就越不容易空。",
            "从这里开始，重点转向把机制翻译成协作方式。",
        ],
        ["src/constants/prompts.ts", "src/query.ts", "src/services/tools/toolExecution.ts", "src/utils/task/framework.ts"],
        "所有高质量技巧的共同点，都是顺着运行时机制，而不是对抗它。",
    )

    add_section_intro(
        prs,
        "第五部分：开发流程与协作建议",
        "不是泛泛给建议，而是从源码约束倒推出更稳的协作方式。",
        [
            "任务怎么写、边界怎么定、验证怎么要求，都会直接影响 Claude Code 的表现。",
            "这些建议不是经验帖，而是前面运行时机制的直接推论。",
            "重点不在“如何让它更聪明”，而在“如何让它更稳定地产出工程结果”。",
        ],
        "源码理解到这里，应该被翻译成一套更有效的协作方式。",
    )
    advice_slides = [
        ("为什么最后要单独讲“开发流程与协作建议”", "前面讲的是机制，这里讲的是怎样顺着机制组织工作。", "这一部分不是经验汇总，而是把 Prompt Stack、Tool Pipeline、Recovery 和 Task Runtime 的约束翻译成团队可执行的工作流。", ["如果停在架构和功能，听众很难把这些理解转化为日常使用方式。", "真正有价值的是：怎样写任务、怎样设边界、怎样做验证、怎样组织长任务。", "源码理解的最终落点，应该是更稳的协作实践。", "因此这部分是全场的落地收束，而不是附录。"], ["src/constants/prompts.ts", "src/services/tools/toolExecution.ts", "src/utils/task/framework.ts"], "前面的运行时机制越清楚，这里的建议越不是“拍脑袋”。"),
        ("建议一：把任务组织成工程任务，而不是聊天请求", "工程化任务表达，是 Claude Code 最自然的输入形式。", "任务至少应该写出目标、范围、约束和验证；这不是格式洁癖，而是顺着 Prompt Stack 的默认角色来组织输入。", ["目标定义它要完成什么。", "范围限制它改到哪里为止。", "约束决定它不能顺手做什么。", "验证决定这次工作如何收口。"], ["src/constants/prompts.ts"], "默认 system prompt 直接把用户请求定义成 software engineering tasks，因此任务越工程化，系统越接近默认最佳状态。", "源码提示词", "The user will primarily request you to perform software engineering tasks."),
        ("建议二：先读代码，再改代码", "显式要求先读关键文件，可以明显减少无关探索和错误改动。", "如果任务涉及真实代码库，就应该先要求它看关键文件、关键函数和现有实现，再决定怎么改。", ["这能降低没读就改、误解上下文和乱造新层的概率。", "对多文件任务尤其重要。", "关键文件点名越清楚，Claude Code 越容易快速进入正确工作面。", "这不是额外约束，而是它的默认偏好。"], ["src/constants/prompts.ts", "src/utils/messages.ts"], "“先读再改”已经写在提示词和 Plan Mode 工作流里，所以显式写出来会更稳。", "源码提示词", "In general, do not propose changes to code you haven't read."),
        ("建议三：最小改动，优先复用", "Claude Code 更适合在边界清楚时做局部、高质量推进。", "明确要求最小必要改动、优先复用现有函数和模式，可以明显抑制顺手重构与过早抽象。", ["这对局部修复、小功能和短周期改动尤其重要。", "系统默认就反对一锤子抽象和无边界优化。", "如果用户不写清楚，它容易把“顺便优化”当成合理延伸。", "越是成熟项目，越应该显式写这一层。"], ["src/constants/prompts.ts", "src/utils/messages.ts"], "最小改动和复用现有模式，是顺着默认 prompt 和搜索工作流在走。", "源码提示词", "Don't create helpers, utilities, or abstractions for one-time operations."),
        ("建议四：专用工具优先，Bash 后置", "可审查、可约束的执行路径，通常比自由 shell 更稳。", "Claude Code 并不是鼓励你默认走 shell；相反，它更偏好 Read、Edit、Glob、Grep 这类 dedicated tools。", ["专用工具更容易被 runtime 审查和记录。", "Bash 更强，但也更高风险。", "复杂执行需求可以把 Bash 留作兜底，而不是默认入口。", "对团队协作场景尤其如此。"], ["src/constants/prompts.ts", "src/tools/BashTool/prompt.ts", "src/tools/BashTool/BashTool.tsx"], "tool pipeline 和 BashTool 的安全层都说明：系统更偏好可审查的执行路径。", "源码提示词", "Do NOT use the Bash tool when a relevant dedicated tool is provided."),
        ("建议五：验证要单独要求，而且结果要如实汇报", "验证不是默认一定会发生，验证结果也不应该靠猜。", "如果任务需要改代码，就应该单独写出要跑什么验证、结果如何汇报、没跑时也要明确说明。", ["这能显著减少“看起来完成了，其实没有真正验证”的情况。", "验证方式越具体，Claude Code 越容易形成闭环。", "对长任务尤其重要，因为它让每个阶段都能有收口点。", "“如实汇报”是默认 prompt 已经明确写出的规则。"], ["src/constants/prompts.ts"], "用户把验证要求显式写出来，其实是在主动利用系统已有的结果约束。", "源码提示词", "Report outcomes faithfully..."),
        ("建议六：高风险动作一定要把确认规则写清楚", "边界越模糊，权限链就越容易放大不确定性。", "涉及删除、覆盖、push、外发、破坏性 git 操作时，最好单独写明确认方式和允许范围。", ["这不是过度谨慎，而是顺着 tool pipeline 和 permission chain 的工作方式。", "对高风险动作，Claude Code 天然会更保守。", "如果用户把授权写得含糊，系统和使用者都会处在不稳状态。", "最好把确认触发条件写成明确规则。"], ["src/constants/prompts.ts", "src/tools/BashTool/prompt.ts", "src/utils/permissions/permissions.ts"], "高风险动作在系统内部是单独对待的，所以用户也必须把这条边界显式表达。", "源码提示词", "Carefully consider the reversibility and blast radius of actions."),
        ("建议七：复杂任务先走 Plan Mode", "先探索、再计划、后实现，比直接开干更稳。", "Plan Mode 的价值在于先快速扫关键文件、形成计划，再问缺失信息或进入执行。", ["对多模块任务、改动范围不确定的任务尤其适合。", "它能帮助系统先建立正确工作面，再开始改。", "也能降低一开始就误判方向的概率。", "Plan Mode 本身已经把这个节奏做成显式流程。"], ["src/utils/messages.ts"], "复杂任务之所以更适合 Plan Mode，不是经验，而是系统里已经存在这条 workflow。"),
        ("建议八：独立查询允许并行，长任务允许分阶段推进", "Claude Code 天生适合把复杂工作拆成可并行、可阶段收口的部分。", "独立查询可以显式允许 parallel，长任务则更适合拆阶段、拆子任务和阶段性收口。", ["并行适合读代码、搜文件、比对多处实现。", "分阶段适合长任务、复杂任务和需要多轮验证的任务。", "这会明显提高输出稳定性，也降低上下文失控。", "task runtime 的存在，本来就在支持这种组织方式。"], ["src/utils/messages.ts", "src/utils/task/framework.ts"], "并行工具调用和 task runtime 都是系统一级能力，所以任务节奏设计会直接影响 Claude Code 的发挥。"),
        ("建议九：哪些任务最适合直接交给 Claude Code", "不是所有任务都适合直接交给它；挑对任务，效果会明显不同。", "最适合的是局部修复、代码理解与定位、需要命令验证的任务，以及可分阶段推进的工程任务。", ["这类任务既需要工具，又需要解释和验证。", "它们边界相对清楚，容易形成阶段性结果。", "也更容易让 Claude Code 的主执行链稳定工作。", "从团队角度看，这类任务最容易形成可复查产出。"], ["src/query.ts", "src/services/tools/toolExecution.ts", "src/utils/task/framework.ts"], "这些任务正好贴合 prompt stack、tool pipeline、task runtime 和 recovery 的优势区间。"),
        ("建议十：哪些任务要谨慎交给 Claude Code", "越依赖隐性边界和模糊背景的任务，越需要先结构化。", "完全模糊的探索、范围极大的全局改动、高风险但边界不清的动作，都不适合直接硬交给 Claude Code。", ["它不是不能做，而是这类任务更容易让工作面漂移。", "此时最好的办法通常不是强行让它“自己想明白”。", "而是先补边界、先拆任务、先确认风险。", "这样才能把任务重新拉回 Claude Code 的优势区间。"], ["src/constants/prompts.ts", "src/services/tools/toolExecution.ts", "src/utils/permissions/permissions.ts"], "不适合直接硬做，不是因为模型弱，而是因为 runtime 对这些任务的稳定工作面本来就更难建立。"),
        ("团队协作时，应该把 Claude Code 当成什么角色", "从团队视角看，Claude Code 更像协作型工程执行体，而不是聊天机器人。", "它最适合承担工程代理、多阶段任务推进器、带工具的验证者，以及文档、代码、验证结果的联合产出器。", ["它能同时做阅读、执行、验证和阶段性结论。", "它适合在明确边界下持续推进。", "它也适合做带证据链的中间产出。", "这比“让它回答问题”更接近它的真实强项。"], ["src/commands.ts", "src/utils/messages.ts", "src/utils/task/framework.ts", "src/QueryEngine.ts"], "commands、Plan Mode、task runtime、structured output 和 tool pipeline 都说明它更适合作为协作型工程执行体。"),
        ("总结：怎样把源码理解真正转化成更好的使用方式", "所有高质量技巧的共同点，不是“更会写 prompt”，而是“更顺着系统工作方式组织任务”。", "先用少量源码理解建立正确模型，再把重心放到产出、使用方式和任务组织上，才是更高价值的 Claude Code 使用方式。", ["先建立最小必要架构模型。", "再理解一次请求如何流过系统。", "再根据产出形态挑任务、定边界、设验证。", "最后把建议落成团队协作方式。"], ["docs/zh/09-Claude Code 源码深度解读-文档.md", "docs/zh/10-Claude Code 源码深度解读-分享提纲.md"], "真正有效的技巧，都可以回到运行时机制解释；这也是这场分享和经验贴的区别。"),
    ]
    for args in advice_slides:
        add_rich_slide(prs.slides.add_slide(prs.slide_layouts[6]), *args)

    prs.save(OUT)
    print(f"saved {OUT}")
    print(f"slides {len(prs.slides)}")


if __name__ == "__main__":
    build_deck()
