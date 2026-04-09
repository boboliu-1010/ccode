from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


TITLE_COLOR = RGBColor(11, 31, 68)
TEXT_COLOR = RGBColor(31, 41, 55)
MUTED_COLOR = RGBColor(71, 85, 105)
ACCENT = RGBColor(32, 82, 149)
ACCENT_2 = RGBColor(221, 234, 251)
BG = RGBColor(248, 250, 252)
LEAD_BG = RGBColor(16, 28, 63)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(219, 228, 240)
FONT = "Arial"


def set_run_font(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_text_frame_style(tf):
    tf.word_wrap = True
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)


def add_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, text, lead=False):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.42), Inches(11.8), Inches(0.75))
    tf = box.text_frame
    set_text_frame_style(tf)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_run_font(r, 26 if not lead else 28, WHITE if lead else TITLE_COLOR, True)
    if not lead:
        line = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.68), Inches(1.18), Inches(2.2), Inches(0.03)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = LINE
        line.line.fill.background()


def add_subtitle(slide, text, lead=False):
    box = slide.shapes.add_textbox(Inches(0.72), Inches(1.4), Inches(11.2), Inches(0.7))
    tf = box.text_frame
    set_text_frame_style(tf)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_run_font(r, 15, WHITE if lead else MUTED_COLOR)


def add_bullets(slide, bullets, top=1.7, left=0.85, width=11.2, height=4.6, color=TEXT_COLOR, size=19):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    set_text_frame_style(tf)
    tf.clear()
    first = True
    for item in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.bullet = True
        p.level = 0
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = item
        set_run_font(r, size, color)


def add_quote(slide, text):
    box = slide.shapes.add_textbox(Inches(0.92), Inches(6.05), Inches(11.0), Inches(0.55))
    tf = box.text_frame
    set_text_frame_style(tf)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_run_font(r, 13, MUTED_COLOR)


def add_section_box(slide, x, y, w, h, title, body):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = ACCENT
    tf = shape.text_frame
    set_text_frame_style(tf)
    tf.clear()
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = title
    set_run_font(r1, 16, ACCENT, True)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(4)
    r2 = p2.add_run()
    r2.text = body
    set_run_font(r2, 13, TEXT_COLOR)


def add_lead_slide(prs, title, subtitle, bullets=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LEAD_BG)
    add_title(slide, title, lead=True)
    add_subtitle(slide, subtitle, lead=True)
    if bullets:
        add_bullets(slide, bullets, top=2.3, left=0.9, width=11.2, height=2.4, color=WHITE, size=18)
    return slide


def add_content_slide(prs, title, bullets, quote=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, title)
    add_bullets(slide, bullets, top=1.65, left=0.85, width=11.2, height=4.8, size=18)
    if quote:
        add_quote(slide, quote)
    return slide


def add_tip_slide(prs, title, tip, why, sources):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title(slide, title)
    add_section_box(slide, 0.8, 1.65, 3.6, 2.0, "推荐写法", tip)
    add_section_box(slide, 4.7, 1.65, 3.8, 2.0, "为什么有效", why)
    add_section_box(slide, 8.8, 1.65, 3.7, 2.0, "源码支撑", sources)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_lead_slide(
        prs,
        "Claude Code 使用技巧与注意事项",
        "从源码反推的用户侧实践建议",
        [
            "先解释 Claude Code 从工程上是什么，再解释怎样用会更稳",
            "每条技巧都附源码支撑，不只讲经验",
        ],
    )
    add_content_slide(
        prs,
        "先给结论：Claude Code 不是什么",
        [
            "它不是自由聊天助手，也不是“随便说一句就稳定帮你把工程活干好”的黑盒",
            "从源码看，它更像 software engineering agent（软件工程代理）",
            "所以用户输入越工程化，Claude Code 的表现通常越稳定",
        ],
        "源码依据：prompts.ts 把默认场景定义为 software engineering task。",
    )
    add_content_slide(
        prs,
        "从源码看：Claude Code 是什么",
        [
            "main.tsx：bootstrap / assembly（启动装配）",
            "QueryEngine.ts：conversation host（会话宿主）",
            "query.ts：turn loop / runtime kernel（轮次循环 / 运行时内核）",
            "toolExecution.ts：tool pipeline（工具执行流水线）",
            "settings / auth / prompt / policy：control plane（控制面）",
        ],
        "这决定了它默认不是把输入当闲聊，而是当软件工程任务。",
    )
    add_content_slide(
        prs,
        "从源码反推：默认偏好的工作方式",
        [
            "先读代码，再改代码",
            "最小必要改动，优先复用已有实现",
            "优先 dedicated tools，而不是默认 Bash",
            "验证后如实汇报",
            "高风险动作默认先确认",
        ],
        "这些偏好主要来自 prompts.ts、messages.ts 和 BashTool prompt。",
    )
    add_tip_slide(
        prs,
        "技巧 1：任务写成“目标 + 范围 + 约束 + 验证”",
        "目标：修复/实现……\n范围：只改……\n约束：不要做无关重构\n验证：跑测试并说明结果",
        "源码把默认场景设成 software engineering task；明确范围和验证能减少分析/实现/顺手优化之间的歧义。",
        "prompts.ts#L221\nprompts.ts#L230\nprompts.ts#L240",
    )
    add_tip_slide(
        prs,
        "技巧 2：明确要求先读代码，再改代码",
        "先看 src/foo.ts 和 src/bar.ts，再决定改法。",
        "系统提示明确偏好先理解现有实现；点名关键文件能减少无关探索和上下文噪音。",
        "prompts.ts#L230\nmessages.ts#L3344",
    )
    add_tip_slide(
        prs,
        "技巧 3：强调最小改动、优先复用",
        "做最小必要改动；优先复用现有函数、工具和模式；不要新造一层抽象。",
        "源码明确反对 one-time helper 和 premature abstraction；这样更容易顺着现有结构修改。",
        "prompts.ts#L200\nprompts.ts#L203\nmessages.ts#L3344",
    )
    add_tip_slide(
        prs,
        "技巧 4：优先 dedicated tools，不要默认 Bash",
        "优先直接读写文件，不要用 Bash 做本可由专用工具完成的事情。",
        "Bash 在源码里是 fallback tool，不是首选；专用工具更容易被审查、解释和权限约束。",
        "prompts.ts#L301\nprompts.ts#L305\nBashTool prompt#L297",
    )
    add_tip_slide(
        prs,
        "技巧 5：验证要求要写清楚，而且要如实汇报",
        "改完后跑相关测试；如果没跑，请明确说明没有跑。",
        "源码对“未验证却宣称完成”是强约束；这样能减少看似完成、其实没验证的风险。",
        "prompts.ts#L211\nprompts.ts#L240",
    )
    add_tip_slide(
        prs,
        "技巧 6：高风险动作要显式要求先确认",
        "涉及 push、删除、覆盖、外部发送、破坏性 git 操作时，先告诉我并确认。",
        "源码对 reversibility 和 blast radius 非常敏感；难以回滚的动作默认就应确认。",
        "prompts.ts#L258\nBashTool prompt#L304",
    )
    add_content_slide(
        prs,
        "两个进阶技巧",
        [
            "独立查询可以显式允许 parallel（并行）",
            "做方案时，先走 Plan Mode 风格：先 Explore，再写计划，再问代码里无法确定的问题",
            "这两条都不是经验贴，而是源码里明确写出来的工作流偏好",
        ],
        "源码依据：prompts.ts#L310、messages.ts#L3336、BashTool prompt#L298。",
    )
    add_content_slide(
        prs,
        "常见误区与注意事项",
        [
            "不要把 Claude Code 当自由聊天助手来用",
            "不要一上来就让它“大改一遍”",
            "不要默认它已经验证过",
            "不要把高风险授权写得太模糊",
        ],
        "这些误区都能在 prompts.ts 的默认约束里找到根源。",
    )
    add_content_slide(
        prs,
        "最后收成 6 条最小使用清单",
        [
            "任务写成：目标 + 范围 + 约束 + 验证",
            "明确要求先读代码",
            "强调最小改动、优先复用",
            "优先 dedicated tools，不要默认走 Bash",
            "验证结果要如实汇报",
            "高风险动作先确认",
        ],
        "这些建议的共同点是：让用户输入尽量贴近 Claude Code 源码里的默认工作方式。",
    )

    out = Path("/Users/bobo/code/claude-code-source-code/docs/zh/09-Claude Code 使用技巧与注意事项-WPS兼容版.pptx")
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
